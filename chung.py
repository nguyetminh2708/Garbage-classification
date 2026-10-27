# -*- coding: utf-8 -*-
"""
Cac hang so va ham dung chung cho cac buoc 01 -> 05.

Thu muc du an:
  Data/anh_tu_chup/<lop>/   anh nhom tu chup, 5 thu muc giay, nhua, huu_co, kim_loai, khac
  Data/realwaste/<lop>/     anh RealWaste da xep thu cong vao 5 thu muc cung ten
  vectors/                  dac trung 512 chieu tu ResNet18 (npz)
  ketqua/                   danh sach anh sau lam sach (csv) va cac bang ket qua (xlsx, csv)
  hinh/                     hinh ve cho bao cao, dat ten theo so hinh trong bao cao
  models/                   mo hinh da huan luyen (joblib)
"""

import copy
import itertools
import time
import warnings
from pathlib import Path

import numpy as np
import pandas as pd
from sklearn.decomposition import PCA
from sklearn.linear_model import SGDClassifier
from sklearn.metrics import (accuracy_score, f1_score, log_loss, hinge_loss,
                             precision_score, recall_score)
from sklearn.model_selection import train_test_split
from sklearn.preprocessing import StandardScaler
from sklearn.svm import SVC

BASE = Path(__file__).parent
DATA_ANH = BASE / "Data"                       # anh goc (khong day len git)
VEC, KQ, HINH, MODEL = BASE / "vectors", BASE / "ketqua", BASE / "hinh", BASE / "models"
for d in (VEC, KQ, HINH, MODEL):
    d.mkdir(exist_ok=True)

CLASSES = ["giay", "nhua", "huu_co", "kim_loai", "khac"]
TEN_LOP = {"giay": "Giấy", "nhua": "Nhựa", "huu_co": "Hữu cơ", "kim_loai": "Kim loại", "khac": "Khác"}
LABELS = np.arange(len(CLASSES))
MODEL_NAMES = ["Softmax Regression", "SVM", "XGBoost"]

SEEDS = [0, 1, 2, 3, 4]            # 5 lan chay, moi lan mot cach chia du lieu
VAL_SIZE, TEST_SIZE = 0.15, 0.15   # train / val / test = 70 / 15 / 15
TOTAL_SAMPLES = 2250               # so anh toi da dung; it hon thi dung het
SAMPLE_SEED = 42
N_COMPONENTS = 128                 # so chieu PCA dung cho thuc nghiem 02 (chon sau thuc nghiem 01)

# ----------------------- luoi sieu tham so (thuc nghiem 02) -----------------------
MAX_EPOCHS, PATIENCE = 200, 15     # Softmax: SGD theo epoch, early stopping theo loss val
XGB_MAX_ROUNDS, XGB_EARLY_STOP = 2000, 50
GRID = {
    "Softmax Regression": {"penalty": ["l2", "l1"], "alpha": [1e-5, 1e-4, 1e-3, 1e-2, 1e-1], "eta0": [1e-4, 1e-3, 1e-2]},
    "SVM": {"kernel": ["linear", "rbf"], "C": [0.1, 1, 10, 100, 1000], "gamma": ["scale", 1e-4, 1e-3, 1e-2, 1e-1]},
    "XGBoost": {"max_depth": [3, 5, 7], "learning_rate": [0.05, 0.1, 0.3], "subsample": [0.7, 1.0], "colsample_bytree": [0.5, 1.0]},
}
# sieu tham so mac dinh (dung cho thuc nghiem 01)
DEFAULT_HP = {
    "Softmax Regression": {"penalty": "l2", "alpha": 1e-3, "eta0": 1e-3},
    "SVM": {"kernel": "rbf", "C": 10, "gamma": "scale"},
    "XGBoost": {"max_depth": 5, "learning_rate": 0.05, "subsample": 0.7, "colsample_bytree": 0.5},
}
LOSS_NAME = {"Softmax Regression": "log loss", "SVM": "hinge loss", "XGBoost": "log loss"}

try:
    import xgboost as xgb
except ImportError:
    xgb = None


# ============================================================================
# Doc dac trung
# ============================================================================
def doc_dac_trung():
    """Doc vectors/features_512.npz; chon toi da TOTAL_SAMPLES anh (phan tang theo lop + nguon).
    Tra ve (meta: DataFrame co cot label, source, y; F: ma tran N x 512)."""
    f = VEC / "features_512.npz"
    if not f.exists():
        raise SystemExit("Thieu vectors/features_512.npz. Chay python 02_trich_dac_trung.py truoc.")
    raw = np.load(f, allow_pickle=True)
    meta = pd.DataFrame({"path": raw["path"], "label": raw["label"], "source": raw["source"]})
    F = raw["feats"]
    meta["y"] = meta.label.map({c: i for i, c in enumerate(CLASSES)})
    if TOTAL_SAMPLES is not None and len(meta) > TOTAL_SAMPLES:
        strat = meta.label + "|" + meta.source
        keep, _ = _strat_split(meta, meta.index, len(meta) - TOTAL_SAMPLES, strat, SAMPLE_SEED)
        meta, F = meta.loc[sorted(keep)].reset_index(drop=True), F[sorted(keep)]
    return meta, F


def _strat_split(meta, index, test_size, strat, seed):
    """Chia phan tang theo (lop, nguon); nhom nao qua it anh thi lui ve phan tang theo lop."""
    try:
        return train_test_split(index, test_size=test_size, stratify=strat, random_state=seed)
    except ValueError:
        warnings.warn("Co nhom (lop, nguon) qua it anh, chi phan tang theo lop.")
        return train_test_split(index, test_size=test_size, stratify=meta.loc[index, "label"], random_state=seed)


def chia_va_pca(meta, F, seed, n_components):
    """Mot lan chay: chia train/val/test roi fit StandardScaler + PCA tren train.
    n_components >= 512 thi khong PCA. Tra ve X, y (dict theo tap), idx, bo tien xu ly, ti le phuong sai giu lai."""
    strat = meta.label + "|" + meta.source
    i_tr, i_tmp = _strat_split(meta, meta.index, VAL_SIZE + TEST_SIZE, strat, seed)
    i_va, i_te = _strat_split(meta, i_tmp, TEST_SIZE / (VAL_SIZE + TEST_SIZE), strat[i_tmp], seed)
    idx = {"train": np.sort(i_tr), "val": np.sort(i_va), "test": np.sort(i_te)}
    s1 = StandardScaler().fit(F[idx["train"]])
    if n_components >= F.shape[1]:
        pca, var_kept = None, 1.0
        Z = {s: s1.transform(F[idx[s]]) for s in idx}
    else:
        pca = PCA(n_components=min(n_components, len(idx["train"])), random_state=seed).fit(s1.transform(F[idx["train"]]))
        var_kept = pca.explained_variance_ratio_.sum()
        Z = {s: pca.transform(s1.transform(F[idx[s]])) for s in idx}
    s2 = StandardScaler().fit(Z["train"])
    X = {s: s2.transform(Z[s]) for s in idx}
    y = {s: meta.y.to_numpy()[idx[s]] for s in idx}
    return X, y, idx, {"scaler_512": s1, "pca": pca, "scaler_pca": s2}, var_kept


# ============================================================================
# Huan luyen tung mo hinh (tra ve: mo hinh, lich su loss, epoch/vong dung)
# ============================================================================
def fit_softmax(X, y, hp, seed):
    rng = np.random.default_rng(seed)
    clf = SGDClassifier(loss="log_loss", penalty=hp["penalty"], alpha=hp["alpha"],
                        learning_rate="constant", eta0=hp["eta0"], random_state=seed)
    h = {"train_loss": [], "val_loss": [], "train_acc": [], "val_acc": []}
    best, best_loss, best_ep, wait = None, np.inf, 0, 0
    for ep in range(1, MAX_EPOCHS + 1):
        p = rng.permutation(len(y["train"]))
        clf.partial_fit(X["train"][p], y["train"][p], classes=LABELS)
        h["train_loss"].append(log_loss(y["train"], clf.predict_proba(X["train"]), labels=LABELS))
        h["val_loss"].append(log_loss(y["val"], clf.predict_proba(X["val"]), labels=LABELS))
        h["train_acc"].append(clf.score(X["train"], y["train"]))
        h["val_acc"].append(clf.score(X["val"], y["val"]))
        if h["val_loss"][-1] < best_loss - 1e-4:
            best, best_loss, best_ep, wait = copy.deepcopy(clf), h["val_loss"][-1], ep, 0
        else:
            wait += 1
            if wait >= PATIENCE:
                break
    return best, h, best_ep


def fit_svm(X, y, hp, seed):
    kw = {"gamma": hp["gamma"]} if hp["kernel"] == "rbf" else {}
    m = SVC(kernel=hp["kernel"], C=hp["C"], random_state=seed, **kw).fit(X["train"], y["train"])
    h = {"train_loss": [hinge_loss(y["train"], m.decision_function(X["train"]), labels=LABELS)],
         "val_loss": [hinge_loss(y["val"], m.decision_function(X["val"]), labels=LABELS)],
         "train_acc": [m.score(X["train"], y["train"])], "val_acc": [m.score(X["val"], y["val"])]}
    return m, h, None


def fit_xgb(X, y, hp, seed):
    if xgb is None:
        raise SystemExit("Thieu thu vien xgboost. Chay: pip install xgboost")
    m = xgb.XGBClassifier(objective="multi:softprob", n_estimators=XGB_MAX_ROUNDS, learning_rate=hp["learning_rate"],
                          max_depth=hp["max_depth"], subsample=hp["subsample"], colsample_bytree=hp["colsample_bytree"],
                          min_child_weight=2, reg_lambda=1.0, eval_metric="mlogloss",
                          early_stopping_rounds=XGB_EARLY_STOP, tree_method="hist", random_state=seed)
    m.fit(X["train"], y["train"], eval_set=[(X["train"], y["train"]), (X["val"], y["val"])], verbose=False)
    ev, bi = m.evals_result(), m.best_iteration
    h = {"train_loss": ev["validation_0"]["mlogloss"], "val_loss": ev["validation_1"]["mlogloss"],
         "train_acc": [m.score(X["train"], y["train"])], "val_acc": [m.score(X["val"], y["val"])]}
    return m, h, bi + 1


FIT = {"Softmax Regression": fit_softmax, "SVM": fit_svm, "XGBoost": fit_xgb}


def hp_str(hp):
    return ", ".join(f"{k}={v}" for k, v in hp.items())


def huan_luyen(name, X, y, hp, seed):
    """Huan luyen 1 mo hinh voi 1 bo sieu tham so. Tra ve (mo hinh, lich su, dong ket qua train/val)."""
    t = time.time()
    m, h, be = FIT[name](X, y, hp, seed)
    k = (be - 1) if be else -1                                    # vi tri diem dung trong lich su loss
    ka = k if len(h["train_acc"]) == len(h["train_loss"]) else -1  # SVM, XGBoost chi ghi accuracy mo hinh cuoi
    r = {"Mô hình": name, "Siêu tham số": hp_str(hp), **hp, "Epoch/vòng dừng": be,
         "Loss train": h["train_loss"][k], "Loss val": h["val_loss"][k],
         "Acc train": h["train_acc"][ka], "Acc val": h["val_acc"][ka], "Thời gian (giây)": time.time() - t}
    return m, h, r


def tieu_chi(name, r):
    """Tieu chi chon tren val (nho hon = tot hon): loss val (Softmax, XGBoost); accuracy val, hoa thi loss val (SVM)."""
    if name == "SVM":
        return (-r["Acc val"], r["Loss val"])
    return (r["Loss val"], -r["Acc val"])


def to_hop_luoi(grid):
    """Moi to hop cua luoi. SVM kernel linear khong dung gamma nen gamma ghi '-' va bo trung."""
    keys, out, seen = list(grid), [], set()
    for vals in itertools.product(*[grid[k] for k in keys]):
        hp = dict(zip(keys, vals))
        if hp.get("kernel") == "linear":
            hp["gamma"] = "-"
        key = tuple(hp.items())
        if key not in seen:
            seen.add(key); out.append(hp)
    return out


def danh_gia_test(m, X, y, src_test):
    """Cac do do tren tap test, ke ca accuracy rieng theo nguon anh."""
    p = m.predict(X["test"])
    out = {"Acc test": accuracy_score(y["test"], p),
           "Precision macro test": precision_score(y["test"], p, average="macro", zero_division=0),
           "Recall macro test": recall_score(y["test"], p, average="macro", zero_division=0),
           "F1 macro test": f1_score(y["test"], p, average="macro", zero_division=0)}
    for src, ten in [("tu_chup", "Acc test tự chụp"), ("realwaste", "Acc test RealWaste")]:
        mask = src_test == src
        out[ten] = accuracy_score(y["test"][mask], p[mask]) if mask.any() else None
    return out, p


def dinh_dang_excel(path):
    """Font, do rong cot, dinh dang so cho cac sheet du lieu (tru sheet dau neu la sheet tong hop)."""
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = load_workbook(path)
    for ws in wb.worksheets:
        if ws.title == "Tong_hop":
            continue
        for row in ws.iter_rows():
            for c in row:
                c.font = Font(name="Arial", size=10, bold=(c.row == 1))
                if c.row == 1:
                    c.fill = PatternFill("solid", fgColor="DCE6F5"); c.alignment = Alignment(wrap_text=True, vertical="center")
        for col in ws.columns:
            head = str(col[0].value or "")
            width = max(len(head), *(len(str(c.value)) if c.value is not None else 0 for c in col[1:200]), 6)
            ws.column_dimensions[col[0].column_letter].width = min(width + 2, 45)
            if any(k in head for k in ["Acc", "Precision", "Recall", "F1", "phương sai", "Loss"]):
                for c in col[1:]: c.number_format = "0.0000"
            if "Thời gian" in head or "dừng TB" in head:
                for c in col[1:]: c.number_format = "0.0"
        ws.freeze_panes = "A2"
    wb.save(path)

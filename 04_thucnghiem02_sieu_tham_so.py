# -*- coding: utf-8 -*-
"""
THUC NGHIEM 02: Khao sat sieu tham so va ket qua chinh (5 lan chay).

Voi N_COMPONENTS chieu (chon sau thuc nghiem 01), moi lan chay:
  1. chia train / val / test 70 / 15 / 15, phan tang theo (lop, nguon anh)
  2. fit StandardScaler + PCA tren train
  3. huan luyen MOI to hop trong luoi sieu tham so cua 3 mo hinh (GRID trong chung.py):
     Softmax 30, SVM 30, XGBoost 36 to hop; mo hinh lap dung early stopping theo loss val
  4. to hop tot nhat tren val duoc danh gia tren test, dung mot lan
Lap 5 lan (seed 0..4): 96 to hop x 5 = 480 luot huan luyen, khoang 1 den 2 gio tren CPU.

Dau ra:
  ketqua/ket_qua_5_lan_chay.xlsx        - 10 sheet (xem bang trong ham xuat_excel)
  models/mo_hinh_5_lan.joblib            - bo tien xu ly va 3 mo hinh cua tung lan chay
  hinh/hinh_4_2_so_sanh_mo_hinh.png      - accuracy va Macro F1 trung binh tren test
  hinh/hinh_4_3_loss_softmax.png         - log loss theo epoch, 5 lan
  hinh/hinh_4_4_svm_validation_curve.png - hinge loss theo C, 5 lan
  hinh/hinh_4_5_loss_xgboost.png         - log loss theo vong boosting, 5 lan
  hinh/hinh_4_6_confusion_matrix.png     - ma tran nham lan cong don 5 lan
"""

import platform
import sys
import time

import joblib
import numpy as np
import pandas as pd
from sklearn.metrics import ConfusionMatrixDisplay, confusion_matrix, precision_recall_fscore_support

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from chung import (CLASSES, GRID, HINH, KQ, LABELS, LOSS_NAME, MAX_EPOCHS, MODEL, MODEL_NAMES, N_COMPONENTS,
                   PATIENCE, SAMPLE_SEED, SEEDS, TEST_SIZE, VAL_SIZE, XGB_EARLY_STOP, XGB_MAX_ROUNDS,
                   chia_va_pca, danh_gia_test, dinh_dang_excel, doc_dac_trung, huan_luyen, tieu_chi, to_hop_luoi, xgb)

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass


# ============================================================================
# 1. Chay thuc nghiem
# ============================================================================
def chay(meta, F):
    n_combo = {n: len(to_hop_luoi(GRID[n])) for n in MODEL_NAMES}
    print("So to hop moi lan chay:", n_combo, "-> tong", sum(n_combo.values()) * len(SEEDS), "luot huan luyen\n")
    detail, class_rows, cm_rows, loss_rows, tune_rows, data_rows = [], [], [], [], [], []
    hists, saved = {}, {}

    for run, seed in enumerate(SEEDS, start=1):
        t_run = time.time()
        X, y, idx, prep, var_kept = chia_va_pca(meta, F, seed, N_COMPONENTS)
        src_test = meta.source.to_numpy()[idx["test"]]
        for s in idx:
            for (lbl, src), n in meta.loc[idx[s]].groupby(["label", "source"]).size().items():
                data_rows.append({"Lần chạy": run, "Seed": seed, "Tập": s, "Lớp": lbl, "Nguồn": src, "Số ảnh": n})

        models = {}
        for name in MODEL_NAMES:
            t_model, best = time.time(), None
            for hp in to_hop_luoi(GRID[name]):
                m, h, r = huan_luyen(name, X, y, hp, seed)
                tune_rows.append({"Lần chạy": run, "Seed": seed, **r})
                if best is None or tieu_chi(name, r) < tieu_chi(name, best[2]):
                    best = (m, h, r)
            m, h, r = best
            models[name] = m; hists[(run, name)] = (h, r["Epoch/vòng dừng"])
            te, p_test = danh_gia_test(m, X, y, src_test)
            detail.append({"Mô hình": name, "Lần chạy": run, "Seed": seed, "Siêu tham số chọn": r["Siêu tham số"],
                           "Số ảnh train": len(y["train"]), "Số ảnh val": len(y["val"]), "Số ảnh test": len(y["test"]),
                           "Số chiều": N_COMPONENTS, "PCA giữ phương sai": var_kept, "Epoch/vòng dừng": r["Epoch/vòng dừng"],
                           "Loại loss": LOSS_NAME[name], "Loss train": r["Loss train"], "Loss val": r["Loss val"],
                           "Acc train": r["Acc train"], "Acc val": r["Acc val"], **te,
                           "Số tổ hợp đã thử": n_combo[name], "Thời gian tìm + huấn luyện (giây)": time.time() - t_model})
            pr, rc, f1, sp = precision_recall_fscore_support(y["test"], p_test, labels=LABELS, zero_division=0)
            for k, c in enumerate(CLASSES):
                class_rows.append({"Mô hình": name, "Lần chạy": run, "Lớp": c, "Precision": pr[k], "Recall": rc[k], "F1": f1[k], "Số ảnh test": sp[k]})
            cm = confusion_matrix(y["test"], p_test, labels=LABELS)
            for k, c in enumerate(CLASSES):
                cm_rows.append({"Mô hình": name, "Lần chạy": run, "Thực tế": c, **{f"Dự đoán {cc}": int(cm[k, j]) for j, cc in enumerate(CLASSES)}})
            if name != "SVM":
                for e in range(len(h["train_loss"])):
                    loss_rows.append({"Mô hình": name, "Lần chạy": run, "Epoch/vòng": e + 1, "Loss train": h["train_loss"][e],
                                      "Loss val": h["val_loss"][e], "Được chọn (early stopping)": "x" if e + 1 == r["Epoch/vòng dừng"] else ""})
        saved[run] = {"prep": prep, "models": models}
        print(f"Lan {run} (seed {seed}, {time.time() - t_run:.0f}s)  acc test  " +
              " | ".join(f"{d['Mô hình']}: {d['Acc test']:.3f} ({d['Siêu tham số chọn']})" for d in detail[-len(MODEL_NAMES):]))

    detail = pd.DataFrame(detail)
    detail["_o"] = detail["Mô hình"].map({n: i for i, n in enumerate(MODEL_NAMES)})
    detail = detail.sort_values(["_o", "Lần chạy"]).drop(columns="_o").reset_index(drop=True)
    return dict(detail=detail, tune=pd.DataFrame(tune_rows), class_rows=pd.DataFrame(class_rows), cm=pd.DataFrame(cm_rows),
                loss=pd.DataFrame(loss_rows), data=pd.DataFrame(data_rows), hists=hists, saved=saved, n_combo=n_combo, meta=meta)


# ============================================================================
# 2. Tong hop sieu tham so
# ============================================================================
def tong_hop_sieu_tham_so(R):
    tune, detail = R["tune"], R["detail"]
    hp = (tune.groupby(["Mô hình", "Siêu tham số"], sort=False)
          .agg(**{"Số lần": ("Acc val", "size"), "Acc train TB": ("Acc train", "mean"), "Acc val TB": ("Acc val", "mean"),
                  "Acc val ĐLC": ("Acc val", "std"), "Loss val TB": ("Loss val", "mean"), "Loss val ĐLC": ("Loss val", "std"),
                  "Epoch/vòng dừng TB": ("Epoch/vòng dừng", "mean"), "Thời gian TB (giây)": ("Thời gian (giây)", "mean")}).reset_index())
    chosen = detail.groupby(["Mô hình", "Siêu tham số chọn"]).size().rename("Số lần được chọn").reset_index().rename(columns={"Siêu tham số chọn": "Siêu tham số"})
    hp = hp.merge(chosen, how="left", on=["Mô hình", "Siêu tham số"]).fillna({"Số lần được chọn": 0})
    hp["Số lần được chọn"] = hp["Số lần được chọn"].astype(int)
    hp["_o"] = hp["Mô hình"].map({n: i for i, n in enumerate(MODEL_NAMES)})
    hp = hp.sort_values(["_o", "Acc val TB"], ascending=[True, False]).drop(columns="_o").reset_index(drop=True)

    # anh huong tung tham so: voi moi gia tri, lay accuracy val TB tot nhat khi cac tham so khac tu do
    for name in MODEL_NAMES:
        for k in GRID[name]:
            if k not in hp: hp[k] = None
        for i in hp.index[hp["Mô hình"] == name]:
            for part in hp.at[i, "Siêu tham số"].split(", "):
                k, v = part.split("="); hp.at[i, k] = v
    marg = []
    for name in MODEL_NAMES:
        sub = hp[hp["Mô hình"] == name]
        for k in GRID[name]:
            for v in [str(x) for x in GRID[name][k]] + (["-"] if k == "gamma" else []):
                ss = sub[sub[k].astype(str) == v]
                if len(ss):
                    marg.append({"Mô hình": name, "Tham số": k, "Giá trị": v, "Số tổ hợp": len(ss),
                                 "Acc val TB tốt nhất": ss["Acc val TB"].max(), "Acc val TB trung bình": ss["Acc val TB"].mean(),
                                 "Acc train TB (tại tổ hợp tốt nhất)": ss.loc[ss["Acc val TB"].idxmax(), "Acc train TB"]})
    cols = ["Mô hình", "Siêu tham số", "Số lần", "Acc train TB", "Acc val TB", "Acc val ĐLC", "Loss val TB", "Loss val ĐLC",
            "Epoch/vòng dừng TB", "Thời gian TB (giây)", "Số lần được chọn"]
    return hp[cols], pd.DataFrame(marg)


# ============================================================================
# 3. Hinh ve
# ============================================================================
def ve_hinh(R):
    detail, hists, tune, cmdf = R["detail"], R["hists"], R["tune"], R["cm"]

    def loss_theo_epoch(name, xlabel, ylabel, fname):
        plt.figure(figsize=(7, 4))
        for r in range(1, len(SEEDS) + 1):
            h, be = hists[(r, name)]
            e = np.arange(1, len(h["train_loss"]) + 1)
            plt.plot(e, h["train_loss"], color="C0", label="train" if r == 1 else None)
            plt.plot(e, h["val_loss"], color="C1", label="val" if r == 1 else None)
            plt.plot(be, h["val_loss"][be - 1], "ko", ms=5, label="điểm dừng (early stopping)" if r == 1 else None)
        plt.xlabel(xlabel); plt.ylabel(ylabel); plt.title(f"{name}: loss train và val qua 5 lần chạy")
        plt.legend(); plt.grid(alpha=0.3); plt.tight_layout(); plt.savefig(HINH / fname, dpi=120); plt.close()

    loss_theo_epoch("Softmax Regression", "Epoch", "Log loss", "hinh_4_3_loss_softmax.png")
    loss_theo_epoch("XGBoost", "Vòng boosting (số cây)", "Log loss", "hinh_4_5_loss_xgboost.png")

    # SVM: validation curve theo C tai gamma duoc chon cua tung lan
    rbf = tune[tune["Mô hình"] == "SVM"].copy(); rbf["gamma"] = rbf["gamma"].astype(str)
    plt.figure(figsize=(7, 4))
    for run in range(1, len(SEEDS) + 1):
        chosen = detail[(detail["Mô hình"] == "SVM") & (detail["Lần chạy"] == run)]["Siêu tham số chọn"].iloc[0]
        g_best = chosen.split("gamma=")[1]
        sub = rbf[(rbf["Lần chạy"] == run) & (rbf.gamma == g_best)].sort_values("C")
        plt.plot(sub.C, sub["Loss train"], "o-", color="C0", label="train" if run == 1 else None)
        plt.plot(sub.C, sub["Loss val"], "o-", color="C1", label="val" if run == 1 else None)
    plt.xscale("log"); plt.xlabel("C"); plt.ylabel("Hinge loss"); plt.title("SVM (kernel RBF): hinge loss theo C qua 5 lần chạy")
    plt.legend(); plt.grid(alpha=0.3); plt.tight_layout(); plt.savefig(HINH / "hinh_4_4_svm_validation_curve.png", dpi=120); plt.close()

    fig, axes = plt.subplots(1, 3, figsize=(15, 4.6))
    for ax, name in zip(axes, MODEL_NAMES):
        cm = cmdf[cmdf["Mô hình"] == name].groupby("Thực tế", sort=False)[[f"Dự đoán {c}" for c in CLASSES]].sum().reindex(CLASSES).to_numpy()
        ConfusionMatrixDisplay(cm, display_labels=CLASSES).plot(ax=ax, colorbar=False, xticks_rotation=30)
        ax.set_title(name); ax.set_xlabel("Dự đoán"); ax.set_ylabel("Thực tế")
    plt.tight_layout(); plt.savefig(HINH / "hinh_4_6_confusion_matrix.png", dpi=120); plt.close()

    g = detail.groupby("Mô hình", sort=False)
    ax = g[["Acc test", "F1 macro test"]].mean().reindex(MODEL_NAMES).plot.bar(figsize=(7, 4), rot=0)
    for cont in ax.containers:
        ax.bar_label(cont, fmt="%.3f", fontsize=9)
    ax.set_ylim(0, 1); ax.set_xlabel(""); ax.set_title("Accuracy và Macro F1 trên tập test (trung bình 5 lần)")
    ax.legend(["Accuracy test", "Macro F1 test"]); ax.grid(axis="y", alpha=0.3)
    plt.tight_layout(); plt.savefig(HINH / "hinh_4_2_so_sanh_mo_hinh.png", dpi=120); plt.close()


# ============================================================================
# 4. Excel
#
# | Sheet | Noi dung |
# | Tong_hop | TB, DLC, min, max cua tung mo hinh qua 5 lan, bang cong thuc Excel tu Chi_tiet_5_lan |
# | Chi_tiet_5_lan | moi dong = 1 mo hinh trong 1 lan chay |
# | Sieu_tham_so_TB | tung to hop: TB +- DLC tren val qua 5 lan, so lan duoc chon |
# | Anh_huong_tung_tham_so | anh huong rieng cua tung sieu tham so |
# | Sieu_tham_so_chi_tiet | moi luot huan luyen |
# | Theo_lop | precision, recall, F1 tung lop, tung lan |
# | Ma_tran_nham_lan | ma tran nham lan tung lan |
# | Loss_theo_epoch | loss train/val tung epoch cua mo hinh duoc chon |
# | Du_lieu_tung_lan | so anh theo tap, lop, nguon cua tung lan chia |
# | Cau_hinh | tham so, phien ban thu vien, may |
# ============================================================================
def xuat_excel(R, hp, marg):
    from openpyxl import load_workbook
    from openpyxl.styles import Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    detail, meta = R["detail"], R["meta"]
    config = pd.DataFrame([
        ("Tổng số ảnh dùng", len(meta)), ("Số lần chạy", len(SEEDS)), ("Seed các lần", ", ".join(map(str, SEEDS))),
        ("Seed chọn ảnh", SAMPLE_SEED), ("Tỉ lệ train / val / test", f"{1-VAL_SIZE-TEST_SIZE:.0%} / {VAL_SIZE:.0%} / {TEST_SIZE:.0%}"),
        ("Trích đặc trưng", "ResNet18 (ImageNet), 512 chiều"), ("Số chiều", f"StandardScaler + PCA {N_COMPONENTS} chiều, fit trên train từng lần"),
        ("Softmax: lưới", str(GRID["Softmax Regression"])), ("Softmax: epoch tối đa / patience", f"{MAX_EPOCHS} / {PATIENCE}"),
        ("SVM: lưới", str(GRID["SVM"])), ("XGBoost: lưới", str(GRID["XGBoost"])),
        ("XGBoost: vòng tối đa / early stopping", f"{XGB_MAX_ROUNDS} / {XGB_EARLY_STOP}"), ("XGBoost: cố định", "min_child_weight 2, reg_lambda 1"),
        ("Số tổ hợp thử mỗi lần chạy", str(R["n_combo"])),
        ("Tiêu chí chọn siêu tham số", "Loss val thấp nhất (Softmax, XGBoost); accuracy val cao nhất, hòa thì loss val thấp hơn (SVM)"),
        ("Ngày chạy", time.strftime("%Y-%m-%d %H:%M")), ("Python", sys.version.split()[0]),
        ("scikit-learn / XGBoost", f"{__import__('sklearn').__version__} / {xgb.__version__ if xgb else 'không có'}"),
        ("Máy", f"{platform.platform()}; CPU {platform.processor() or 'không rõ'}"),
    ], columns=["Tham số", "Giá trị"])
    tune_cols = ["Lần chạy", "Seed", "Mô hình", "Siêu tham số", "Epoch/vòng dừng", "Loss train", "Loss val", "Acc train", "Acc val", "Thời gian (giây)"]

    xlsx = KQ / "ket_qua_5_lan_chay.xlsx"
    with pd.ExcelWriter(xlsx, engine="openpyxl") as w:
        pd.DataFrame().to_excel(w, sheet_name="Tong_hop")
        detail.to_excel(w, sheet_name="Chi_tiet_5_lan", index=False)
        hp.to_excel(w, sheet_name="Sieu_tham_so_TB", index=False)
        marg.to_excel(w, sheet_name="Anh_huong_tung_tham_so", index=False)
        R["tune"][tune_cols].to_excel(w, sheet_name="Sieu_tham_so_chi_tiet", index=False)
        R["class_rows"].to_excel(w, sheet_name="Theo_lop", index=False)
        R["cm"].to_excel(w, sheet_name="Ma_tran_nham_lan", index=False)
        R["loss"].to_excel(w, sheet_name="Loss_theo_epoch", index=False)
        R["data"].to_excel(w, sheet_name="Du_lieu_tung_lan", index=False)
        config.to_excel(w, sheet_name="Cau_hinh", index=False)
    dinh_dang_excel(xlsx)

    # sheet Tong_hop: cong thuc tham chieu Chi_tiet_5_lan de bam vao o la thay nguon
    wb = load_workbook(xlsx); ws, src = wb["Tong_hop"], wb["Chi_tiet_5_lan"]
    FONT, fill, thin = "Arial", PatternFill("solid", fgColor="DCE6F5"), Side(style="thin", color="BFBFBF")
    col_of = {src.cell(1, j).value: get_column_letter(j) for j in range(1, src.max_column + 1)}
    ws["A1"] = "TỔNG HỢP KẾT QUẢ 5 LẦN CHẠY"; ws["A1"].font = Font(name=FONT, size=13, bold=True)
    ws["A2"] = f"{len(meta)} ảnh, 5 lớp, {N_COMPONENTS} chiều; mỗi lần chia lại train/val/test và chọn lại siêu tham số trên val. Mọi số trong bảng là công thức tính từ sheet Chi_tiet_5_lan."
    ws["A2"].font = Font(name=FONT, size=10, italic=True, color="555555")
    METRICS = ["Acc train", "Acc val", "Acc test", "Precision macro test", "Recall macro test", "F1 macro test",
               "Acc test tự chụp", "Acc test RealWaste", "Loss train", "Loss val", "Thời gian tìm + huấn luyện (giây)"]
    STATS = [("Trung bình", "AVERAGE"), ("Độ lệch chuẩn", "STDEV"), ("Min", "MIN"), ("Max", "MAX")]
    r = 4
    for name in MODEL_NAMES:
        rows = detail.index[detail["Mô hình"] == name] + 2
        r0, r1 = rows.min(), rows.max()
        ws.cell(r, 1, name).font = Font(name=FONT, size=11, bold=True)
        ws.cell(r, 2, f"(dòng {r0} đến {r1} của Chi_tiet_5_lan)").font = Font(name=FONT, size=9, italic=True, color="777777"); r += 1
        ws.cell(r, 1, "Chỉ số")
        for j, (lbl, _) in enumerate(STATS): ws.cell(r, 2 + j, lbl)
        ws.cell(r, 2 + len(STATS), "Viết vào báo cáo (TB ± ĐLC)")
        for j in range(1, 3 + len(STATS)):
            c = ws.cell(r, j); c.font = Font(name=FONT, size=10, bold=True); c.fill = fill; c.border = Border(bottom=thin)
        r += 1
        for met in METRICS:
            rng_ = f"Chi_tiet_5_lan!{col_of[met]}{r0}:{col_of[met]}{r1}"
            ws.cell(r, 1, met).font = Font(name=FONT, size=10)
            for j, (_, fn) in enumerate(STATS):
                c = ws.cell(r, 2 + j, f"={fn}({rng_})"); c.font = Font(name=FONT, size=10); c.number_format = "0.0" if "Thời gian" in met else "0.0000"
            nd = 1 if "Thời gian" in met else 4
            ws.cell(r, 2 + len(STATS), f'=FIXED(AVERAGE({rng_}),{nd},TRUE)&" ± "&FIXED(STDEV({rng_}),{nd},TRUE)').font = Font(name=FONT, size=10)
            r += 1
        r += 1
    ws.cell(r, 1, "Accuracy test và siêu tham số được chọn ở từng lần chạy").font = Font(name=FONT, size=11, bold=True); r += 1
    ws.cell(r, 1, "Mô hình")
    for k in range(len(SEEDS)): ws.cell(r, 2 + k, f"Lần {k + 1} (seed {SEEDS[k]})")
    ws.cell(r, 2 + len(SEEDS), "Trung bình")
    for j in range(1, 3 + len(SEEDS)):
        c = ws.cell(r, j); c.font = Font(name=FONT, size=10, bold=True); c.fill = fill
    r += 1
    for name in MODEL_NAMES:
        rows = detail.index[detail["Mô hình"] == name] + 2
        ws.cell(r, 1, name).font = Font(name=FONT, size=10)
        for k, rr in enumerate(rows):
            c = ws.cell(r, 2 + k, f"=Chi_tiet_5_lan!{col_of['Acc test']}{rr}"); c.number_format = "0.0000"; c.font = Font(name=FONT, size=10)
            ws.cell(r + 1, 2 + k, f"=Chi_tiet_5_lan!{col_of['Siêu tham số chọn']}{rr}").font = Font(name=FONT, size=8, italic=True, color="555555")
        c = ws.cell(r, 2 + len(SEEDS), f"=AVERAGE({get_column_letter(2)}{r}:{get_column_letter(1 + len(SEEDS))}{r})")
        c.number_format = "0.0000"; c.font = Font(name=FONT, size=10, bold=True); r += 2
    ws.column_dimensions["A"].width = 26
    for j in range(2, 3 + max(len(STATS), len(SEEDS))): ws.column_dimensions[get_column_letter(j)].width = 22
    wb.save(xlsx)


def main() -> None:
    meta, F = doc_dac_trung()
    print(f"{len(meta)} anh, dac trung {F.shape[1]} chieu, PCA {N_COMPONENTS} chieu")
    R = chay(meta, F)
    hp, marg = tong_hop_sieu_tham_so(R)

    g = R["detail"].groupby("Mô hình", sort=False)
    print("\n=== Ket qua trung binh 5 lan (TB +- DLC) ===")
    for met in ["Acc train", "Acc val", "Acc test", "F1 macro test", "Acc test tự chụp", "Acc test RealWaste"]:
        print(f"  {met:20s}" + "  ".join(f"{n}: {g[met].mean()[n]:.3f} ± {g[met].std()[n]:.3f}" for n in MODEL_NAMES))

    ve_hinh(R)
    xuat_excel(R, hp, marg)
    joblib.dump({"runs": R["saved"], "classes": CLASSES, "seeds": SEEDS, "n_components": N_COMPONENTS}, MODEL / "mo_hinh_5_lan.joblib")
    print("\nDa ghi: ketqua/ket_qua_5_lan_chay.xlsx, models/mo_hinh_5_lan.joblib, hinh/hinh_4_2 den hinh_4_6")


if __name__ == "__main__":
    main()

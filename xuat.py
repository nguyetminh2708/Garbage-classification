# -*- coding: utf-8 -*-
"""
Phan xuat ket qua: tinh bang tong hop, ghi Excel / CSV va ve hinh cho cac buoc 01 -> 04.
Tach rieng de cac file buoc chi con logic xu ly. Ten ham noi ro no phuc vu buoc nao va tao ra file nao.

  01: thong_ke_du_lieu, xuat_thong_ke_du_lieu, ve_hinh_du_lieu     -> ketqua/thong_ke_du_lieu.xlsx, danh_sach_anh.csv, hinh_3_1, hinh_3_2
  02: ve_hinh_truoc_sau, ve_hinh_pca                                -> hinh_3_3, hinh_3_4a, hinh_3_4b
  03: tong_hop_pca, xuat_thucnghiem01, ve_hinh_thucnghiem01         -> ketqua/thucnghiem01_pca*.csv, hinh_4_1
  04: tong_hop_sieu_tham_so, ve_hinh_thucnghiem02, bang_cau_hinh,
      ghi_cac_sheet, sheet_tong_hop                                 -> ketqua/ket_qua_5_lan_chay.xlsx, hinh_4_2 den hinh_4_6
"""

import platform
import sys
import time
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sklearn.metrics import ConfusionMatrixDisplay

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from chung import (CLASSES, DATA_ANH, GRID, HINH, KQ, MAX_EPOCHS, MODEL_NAMES, N_COMPONENTS, PATIENCE,
                   SAMPLE_SEED, SEEDS, TEST_SIZE, VAL_SIZE, XGB_EARLY_STOP, XGB_MAX_ROUNDS, chia_va_pca, xgb)


def rel(p: str) -> str:
    """Duong dan tuong doi so voi thu muc Data, de ghi vao Excel cho gon."""
    try:
        return str(Path(p).relative_to(DATA_ANH))
    except ValueError:
        return p


# ============================================================================
# Dung chung
# ============================================================================
def dinh_dang_excel(path):
    """Font, do rong cot, dinh dang so cho cac sheet du lieu (tru sheet dau neu la sheet tong hop)."""
    wb = load_workbook(path)
    for ws in wb.worksheets:
        if ws.title == "Tong_hop":
            continue
        for row in ws.iter_rows():
            for c in row:
                c.font = Font(name="Arial", size=10, bold=(c.row == 1))
                if c.row == 1:
                    c.fill = PatternFill("solid", fgColor="DCE6F5"); c.alignment = Alignment(wrap_text=True, vertical="center")
        for col in ws.columns:
            head = str(col[0].value or "")
            width = max(len(head), *(len(str(c.value)) if c.value is not None else 0 for c in col[1:200]), 6)
            ws.column_dimensions[col[0].column_letter].width = min(width + 2, 45)
            if any(k in head for k in ["Acc", "Precision", "Recall", "F1", "phương sai", "Loss"]):
                for c in col[1:]: c.number_format = "0.0000"
            if "Thời gian" in head or "dừng TB" in head:
                for c in col[1:]: c.number_format = "0.0"
        ws.freeze_panes = "A2"
    wb.save(path)


# ============================================================================
# Buoc 01: thong ke du lieu
# ============================================================================
def thong_ke_du_lieu(df: pd.DataFrame, n_truoc: pd.DataFrame, own_avail: dict, rw_avail: dict, min_side: int) -> dict:
    """Tra ve cac bang thong ke, moi bang la mot sheet Excel."""
    n_sau = pd.crosstab(df.label, df.source).reindex(CLASSES).fillna(0).astype(int)
    tong_quan = pd.DataFrame({"Lớp": CLASSES})
    for src in ["tu_chup", "realwaste"]:
        tong_quan[f"{src}: trước làm sạch"] = n_truoc[src].to_numpy() if src in n_truoc else 0
        tong_quan[f"{src}: sau làm sạch"] = n_sau[src].to_numpy() if src in n_sau else 0
    tong_quan["Tổng sau làm sạch"] = n_sau.sum(axis=1).to_numpy()
    tong_quan["Ảnh tự chụp có sẵn"] = [own_avail[c] for c in CLASSES]
    tong_quan["Ảnh RealWaste có sẵn"] = [rw_avail[c] for c in CLASSES]
    tong_quan.loc[len(tong_quan)] = ["Tổng"] + [tong_quan[c].sum() for c in tong_quan.columns[1:]]

    kich_thuoc = (df.groupby(["label", "source"])
                  .agg(**{"Số ảnh": ("path", "size"), "Rộng min": ("w", "min"), "Rộng trung vị": ("w", "median"), "Rộng max": ("w", "max"),
                          "Cao min": ("h", "min"), "Cao trung vị": ("h", "median"), "Cao max": ("h", "max"),
                          "Tỉ lệ ảnh vuông": ("square", "mean"), f"Ảnh có cạnh < {min_side}": ("small", "sum")})
                  .reset_index().rename(columns={"label": "Lớp", "source": "Nguồn"}))
    dinh_dang = pd.crosstab([df.source, df.ext], df.label).reindex(columns=CLASSES).fillna(0).astype(int).reset_index().rename(columns={"source": "Nguồn", "ext": "Định dạng"})
    che_do_mau = pd.crosstab(df.source, df["mode"]).reset_index().rename(columns={"source": "Nguồn"})
    loai_con = df.groupby(["label", "source", "orig_class"]).size().rename("Số ảnh").reset_index().rename(columns={"label": "Lớp", "source": "Nguồn", "orig_class": "Loại con"})
    danh_sach = df[["path", "label", "source", "orig_class", "w", "h", "ext", "mode"]].assign(path=df.path.map(rel)).rename(
        columns={"path": "Đường dẫn", "label": "Lớp", "source": "Nguồn", "orig_class": "Loại con", "w": "Rộng", "h": "Cao", "ext": "Định dạng", "mode": "Chế độ màu"})
    return {"tong_quan": tong_quan, "kich_thuoc": kich_thuoc, "dinh_dang": dinh_dang, "che_do_mau": che_do_mau, "loai_con": loai_con, "danh_sach": danh_sach}


def xuat_thong_ke_du_lieu(df: pd.DataFrame, tk: dict, loai_bo: list[dict]) -> None:
    """ketqua/thong_ke_du_lieu.xlsx (6 sheet) va ketqua/danh_sach_anh.csv cho buoc 02."""
    xlsx = KQ / "thong_ke_du_lieu.xlsx"
    with pd.ExcelWriter(xlsx, engine="openpyxl") as w:
        tk["tong_quan"].to_excel(w, sheet_name="Tong_quan", index=False)
        pd.DataFrame(loai_bo, columns=["Lý do", "Đường dẫn", "Lớp", "Nguồn", "Ghi chú"]).to_excel(w, sheet_name="Loai_bo", index=False)
        tk["kich_thuoc"].to_excel(w, sheet_name="Kich_thuoc", index=False)
        tk["dinh_dang"].to_excel(w, sheet_name="Dinh_dang", index=False)
        tk["che_do_mau"].to_excel(w, sheet_name="Dinh_dang", index=False, startrow=len(tk["dinh_dang"]) + 3)
        tk["loai_con"].to_excel(w, sheet_name="Loai_con", index=False)
        tk["danh_sach"].to_excel(w, sheet_name="Danh_sach_anh", index=False)
    dinh_dang_excel(xlsx)
    df.drop(columns=["ok", "md5", "err"]).to_csv(KQ / "danh_sach_anh.csv", index=False, encoding="utf-8-sig")


def ve_hinh_du_lieu(df: pd.DataFrame, min_side: int) -> None:
    """hinh_3_1: so anh moi lop theo nguon; hinh_3_2: phan bo canh ngan cua anh."""
    counts = pd.crosstab(df.label, df.source).reindex(CLASSES).fillna(0)
    ax = counts.plot.bar(stacked=True, figsize=(7, 4), rot=0)
    ax.set_xlabel(""); ax.set_ylabel("Số ảnh"); ax.set_title("Số ảnh mỗi lớp theo nguồn"); ax.grid(axis="y", alpha=0.3)
    plt.tight_layout(); plt.savefig(HINH / "hinh_3_1_phan_bo_lop.png", dpi=120); plt.close()

    plt.figure(figsize=(7, 4))
    for src, sub in df.groupby("source"):
        plt.hist(sub[["w", "h"]].min(axis=1), bins=30, alpha=0.7, label=f"{src} ({len(sub)} ảnh)")
    plt.axvline(min_side, color="red", ls="--", label=f"ngưỡng {min_side}px")
    plt.xlabel("Cạnh ngắn của ảnh (điểm ảnh)"); plt.ylabel("Số ảnh"); plt.title("Phân bố kích thước ảnh theo nguồn")
    plt.legend(); plt.grid(alpha=0.3)
    plt.tight_layout(); plt.savefig(HINH / "hinh_3_2_phan_bo_kich_thuoc.png", dpi=120); plt.close()


# ============================================================================
# Buoc 02: hinh minh hoa tien xu ly va PCA
# ============================================================================
def ve_hinh_truoc_sau(df: pd.DataFrame, load_rgb, to_tensor, mean, std, img_size: int) -> None:
    """hinh_3_3: moi lop mot anh, hang tren la anh goc, hang duoi la anh sau chuan hoa."""
    def denorm(t): return (t.permute(1, 2, 0).numpy() * std + mean).clip(0, 1)
    sample = df.groupby("label").sample(1, random_state=42).set_index("label").reindex(CLASSES)
    fig, axes = plt.subplots(2, len(CLASSES), figsize=(3 * len(CLASSES), 6.4))
    for j, (label, r) in enumerate(sample.iterrows()):
        axes[0, j].imshow(load_rgb(r.path)); axes[0, j].set_title(f"{label}\n{r.w}x{r.h} ({r.source})", fontsize=9)
        axes[1, j].imshow(denorm(to_tensor(load_rgb(r.path)))); axes[1, j].set_title(f"sau xử lý {img_size}x{img_size}", fontsize=9)
    for a in axes.ravel(): a.axis("off")
    plt.tight_layout(); plt.savefig(HINH / "hinh_3_3_anh_truoc_sau.png", dpi=120); plt.close()


def ve_hinh_pca(df: pd.DataFrame, feats: np.ndarray) -> float:
    """hinh_3_4a: phuong sai tich luy theo so chieu; hinh_3_4b: du lieu tren 2 thanh phan chinh dau.
    Chi de minh hoa: fit PCA tren tap train cua mot lan chia (seed 42). Tra ve ti le phuong sai giu lai."""
    meta = df[["path", "label", "source"]].copy(); meta["y"] = meta.label.map({c: i for i, c in enumerate(CLASSES)})
    _, _, _, prep, kept = chia_va_pca(meta, feats, 42, N_COMPONENTS)
    cum = np.cumsum(prep["pca"].explained_variance_ratio_)
    plt.figure(figsize=(7, 4))
    plt.plot(np.arange(1, len(cum) + 1), cum)
    plt.axhline(kept, color="red", ls="--", label=f"{len(cum)} chiều giữ {kept:.1%}")
    plt.xlabel("Số thành phần chính"); plt.ylabel("Phương sai tích lũy được giữ lại")
    plt.title("PCA: phương sai giữ lại theo số chiều"); plt.ylim(0, 1.02); plt.legend(); plt.grid(alpha=0.3)
    plt.tight_layout(); plt.savefig(HINH / "hinh_3_4a_pca_phuong_sai.png", dpi=120); plt.close()

    Z = prep["pca"].transform(prep["scaler_512"].transform(feats))
    plt.figure(figsize=(7, 5.5))
    for i, c in enumerate(CLASSES):
        pts = Z[meta.y.to_numpy() == i]
        plt.scatter(pts[:, 0], pts[:, 1], s=12, alpha=0.6, label=c)
    plt.xlabel("PC1"); plt.ylabel("PC2"); plt.title("Dữ liệu trên 2 thành phần chính đầu"); plt.legend(); plt.grid(alpha=0.3)
    plt.tight_layout(); plt.savefig(HINH / "hinh_3_4b_pca_2d.png", dpi=120); plt.close()
    return kept


# ============================================================================
# Buoc 03: thuc nghiem 01
# ============================================================================
def tong_hop_pca(df: pd.DataFrame) -> pd.DataFrame:
    """Trung binh va do lech chuan theo (mo hinh, so chieu)."""
    return (df.groupby(["Mô hình", "Số chiều"], sort=False)
            .agg(**{"PCA giữ phương sai": ("PCA giữ phương sai", "mean"), "Acc train TB": ("Acc train", "mean"),
                    "Acc val TB": ("Acc val", "mean"), "Acc val ĐLC": ("Acc val", "std"),
                    "Acc test TB": ("Acc test", "mean"), "Acc test ĐLC": ("Acc test", "std"),
                    "F1 test TB": ("F1 macro test", "mean"), "Thời gian TB (giây)": ("Thời gian (giây)", "mean")}).reset_index())


def xuat_thucnghiem01(df: pd.DataFrame, tb: pd.DataFrame) -> None:
    df.to_csv(KQ / "thucnghiem01_pca.csv", index=False, encoding="utf-8-sig")
    tb.to_csv(KQ / "thucnghiem01_pca_tb.csv", index=False, encoding="utf-8-sig", float_format="%.4f")


def ve_hinh_thucnghiem01(df: pd.DataFrame, pca_dims: list) -> None:
    """hinh_4_1: accuracy val trung binh theo so chieu, moi mo hinh mot duong."""
    plt.figure(figsize=(7, 4))
    for name in MODEL_NAMES:
        mu = df[df["Mô hình"] == name].groupby("Số chiều")["Acc val"].mean()
        plt.plot(mu.index, mu.values, marker="o", label=name)
    plt.xscale("log", base=2); plt.xticks(pca_dims, [str(d) for d in pca_dims])
    plt.xlabel("Số chiều sau PCA (512 = không giảm chiều)"); plt.ylabel("Accuracy val (trung bình 5 lần)")
    plt.title("Khảo sát số chiều PCA"); plt.legend(); plt.grid(alpha=0.3)
    plt.tight_layout(); plt.savefig(HINH / "hinh_4_1_khao_sat_pca.png", dpi=120); plt.close()


# ============================================================================
# Buoc 04: thuc nghiem 02
# ============================================================================
def tong_hop_sieu_tham_so(R):
    tune, detail = R["tune"], R["detail"]
    hp = (tune.groupby(["Mô hình", "Siêu tham số"], sort=False)
          .agg(**{"Số lần": ("Acc val", "size"), "Acc train TB": ("Acc train", "mean"), "Acc val TB": ("Acc val", "mean"),
                  "Acc val ĐLC": ("Acc val", "std"), "Loss val TB": ("Loss val", "mean"), "Loss val ĐLC": ("Loss val", "std"),
                  "Epoch/vòng dừng TB": ("Epoch/vòng dừng", "mean"), "Thời gian TB (giây)": ("Thời gian (giây)", "mean")}).reset_index())
    chosen = detail.groupby(["Mô hình", "Siêu tham số chọn"]).size().rename("Số lần được chọn").reset_index().rename(columns={"Siêu tham số chọn": "Siêu tham số"})
    hp = hp.merge(chosen, how="left", on=["Mô hình", "Siêu tham số"]).fillna({"Số lần được chọn": 0})
    hp["Số lần được chọn"] = hp["Số lần được chọn"].astype(int)
    hp["_o"] = hp["Mô hình"].map({n: i for i, n in enumerate(MODEL_NAMES)})
    hp = hp.sort_values(["_o", "Acc val TB"], ascending=[True, False]).drop(columns="_o").reset_index(drop=True)

    # anh huong tung tham so: voi moi gia tri, lay accuracy val TB tot nhat khi cac tham so khac tu do
    for name in MODEL_NAMES:
        for k in GRID[name]:
            if k not in hp: hp[k] = None
        for i in hp.index[hp["Mô hình"] == name]:
            for part in hp.at[i, "Siêu tham số"].split(", "):
                k, v = part.split("="); hp.at[i, k] = v
    marg = []
    for name in MODEL_NAMES:
        sub = hp[hp["Mô hình"] == name]
        for k in GRID[name]:
            for v in [str(x) for x in GRID[name][k]] + (["-"] if k == "gamma" else []):
                ss = sub[sub[k].astype(str) == v]
                if len(ss):
                    marg.append({"Mô hình": name, "Tham số": k, "Giá trị": v, "Số tổ hợp": len(ss),
                                 "Acc val TB tốt nhất": ss["Acc val TB"].max(), "Acc val TB trung bình": ss["Acc val TB"].mean(),
                                 "Acc train TB (tại tổ hợp tốt nhất)": ss.loc[ss["Acc val TB"].idxmax(), "Acc train TB"]})
    cols = ["Mô hình", "Siêu tham số", "Số lần", "Acc train TB", "Acc val TB", "Acc val ĐLC", "Loss val TB", "Loss val ĐLC",
            "Epoch/vòng dừng TB", "Thời gian TB (giây)", "Số lần được chọn"]
    return hp[cols], pd.DataFrame(marg)


def ve_hinh_thucnghiem02(R):
    detail, hists, tune, cmdf = R["detail"], R["hists"], R["tune"], R["cm"]

    def loss_theo_epoch(name, xlabel, ylabel, fname):
        plt.figure(figsize=(7, 4))
        for r in range(1, len(SEEDS) + 1):
            h, be = hists[(r, name)]
            e = np.arange(1, len(h["train_loss"]) + 1)
            plt.plot(e, h["train_loss"], color="C0", label="train" if r == 1 else None)
            plt.plot(e, h["val_loss"], color="C1", label="val" if r == 1 else None)
            plt.plot(be, h["val_loss"][be - 1], "ko", ms=5, label="điểm dừng (early stopping)" if r == 1 else None)
        plt.xlabel(xlabel); plt.ylabel(ylabel); plt.title(f"{name}: loss train và val qua 5 lần chạy")
        plt.legend(); plt.grid(alpha=0.3); plt.tight_layout(); plt.savefig(HINH / fname, dpi=120); plt.close()

    loss_theo_epoch("Softmax Regression", "Epoch", "Log loss", "hinh_4_3_loss_softmax.png")
    loss_theo_epoch("XGBoost", "Vòng boosting (số cây)", "Log loss", "hinh_4_5_loss_xgboost.png")

    # SVM: validation curve theo C tai gamma duoc chon cua tung lan
    rbf = tune[tune["Mô hình"] == "SVM"].copy(); rbf["gamma"] = rbf["gamma"].astype(str)
    plt.figure(figsize=(7, 4))
    for run in range(1, len(SEEDS) + 1):
        chosen = detail[(detail["Mô hình"] == "SVM") & (detail["Lần chạy"] == run)]["Siêu tham số chọn"].iloc[0]
        g_best = chosen.split("gamma=")[1]
        sub = rbf[(rbf["Lần chạy"] == run) & (rbf.gamma == g_best)].sort_values("C")
        plt.plot(sub.C, sub["Loss train"], "o-", color="C0", label="train" if run == 1 else None)
        plt.plot(sub.C, sub["Loss val"], "o-", color="C1", label="val" if run == 1 else None)
    plt.xscale("log"); plt.xlabel("C"); plt.ylabel("Hinge loss"); plt.title("SVM (kernel RBF): hinge loss theo C qua 5 lần chạy")
    plt.legend(); plt.grid(alpha=0.3); plt.tight_layout(); plt.savefig(HINH / "hinh_4_4_svm_validation_curve.png", dpi=120); plt.close()

    fig, axes = plt.subplots(1, 3, figsize=(15, 4.6))
    for ax, name in zip(axes, MODEL_NAMES):
        cm = cmdf[cmdf["Mô hình"] == name].groupby("Thực tế", sort=False)[[f"Dự đoán {c}" for c in CLASSES]].sum().reindex(CLASSES).to_numpy()
        ConfusionMatrixDisplay(cm, display_labels=CLASSES).plot(ax=ax, colorbar=False, xticks_rotation=30)
        ax.set_title(name); ax.set_xlabel("Dự đoán"); ax.set_ylabel("Thực tế")
    plt.tight_layout(); plt.savefig(HINH / "hinh_4_6_confusion_matrix.png", dpi=120); plt.close()

    g = detail.groupby("Mô hình", sort=False)
    ax = g[["Acc test", "F1 macro test"]].mean().reindex(MODEL_NAMES).plot.bar(figsize=(7, 4), rot=0)
    for cont in ax.containers:
        ax.bar_label(cont, fmt="%.3f", fontsize=9)
    ax.set_ylim(0, 1); ax.set_xlabel(""); ax.set_title("Accuracy và Macro F1 trên tập test (trung bình 5 lần)")
    ax.legend(["Accuracy test", "Macro F1 test"]); ax.grid(axis="y", alpha=0.3)
    plt.tight_layout(); plt.savefig(HINH / "hinh_4_2_so_sanh_mo_hinh.png", dpi=120); plt.close()


def bang_cau_hinh(R) -> pd.DataFrame:
    """Sheet Cau_hinh: tham so da dung, phien ban thu vien, may."""
    meta = R["meta"]
    return pd.DataFrame([
        ("Tổng số ảnh dùng", len(meta)), ("Số lần chạy", len(SEEDS)), ("Seed các lần", ", ".join(map(str, SEEDS))),
        ("Seed chọn ảnh", SAMPLE_SEED), ("Tỉ lệ train / val / test", f"{1-VAL_SIZE-TEST_SIZE:.0%} / {VAL_SIZE:.0%} / {TEST_SIZE:.0%}"),
        ("Trích đặc trưng", "ResNet18 (ImageNet), 512 chiều"), ("Số chiều", f"StandardScaler + PCA {N_COMPONENTS} chiều, fit trên train từng lần"),
        ("Softmax: lưới", str(GRID["Softmax Regression"])), ("Softmax: epoch tối đa / patience", f"{MAX_EPOCHS} / {PATIENCE}"),
        ("SVM: lưới", str(GRID["SVM"])), ("XGBoost: lưới", str(GRID["XGBoost"])),
        ("XGBoost: vòng tối đa / early stopping", f"{XGB_MAX_ROUNDS} / {XGB_EARLY_STOP}"), ("XGBoost: cố định", "min_child_weight 2, reg_lambda 1"),
        ("Số tổ hợp thử mỗi lần chạy", str(R["n_combo"])),
        ("Tiêu chí chọn siêu tham số", "Loss val thấp nhất (Softmax, XGBoost); accuracy val cao nhất, hòa thì loss val thấp hơn (SVM)"),
        ("Ngày chạy", time.strftime("%Y-%m-%d %H:%M")), ("Python", sys.version.split()[0]),
        ("scikit-learn / XGBoost", f"{__import__('sklearn').__version__} / {xgb.__version__ if xgb else 'không có'}"),
        ("Máy", f"{platform.platform()}; CPU {platform.processor() or 'không rõ'}"),
    ], columns=["Tham số", "Giá trị"])


def ghi_cac_sheet(R, hp, marg, config):
    """Ghi 10 sheet so lieu vao ketqua/ket_qua_5_lan_chay.xlsx (sheet Tong_hop de trong, dien o ham sau)."""
    tune_cols = ["Lần chạy", "Seed", "Mô hình", "Siêu tham số", "Epoch/vòng dừng", "Loss train", "Loss val", "Acc train", "Acc val", "Thời gian (giây)"]

    xlsx = KQ / "ket_qua_5_lan_chay.xlsx"
    with pd.ExcelWriter(xlsx, engine="openpyxl") as w:
        pd.DataFrame().to_excel(w, sheet_name="Tong_hop")
        R["detail"].to_excel(w, sheet_name="Chi_tiet_5_lan", index=False)
        hp.to_excel(w, sheet_name="Sieu_tham_so_TB", index=False)
        marg.to_excel(w, sheet_name="Anh_huong_tung_tham_so", index=False)
        R["tune"][tune_cols].to_excel(w, sheet_name="Sieu_tham_so_chi_tiet", index=False)
        R["class_rows"].to_excel(w, sheet_name="Theo_lop", index=False)
        R["cm"].to_excel(w, sheet_name="Ma_tran_nham_lan", index=False)
        R["loss"].to_excel(w, sheet_name="Loss_theo_epoch", index=False)
        R["data"].to_excel(w, sheet_name="Du_lieu_tung_lan", index=False)
        config.to_excel(w, sheet_name="Cau_hinh", index=False)
    dinh_dang_excel(xlsx)
    return xlsx


def sheet_tong_hop(xlsx, R) -> None:
    """Sheet Tong_hop: TB, DLC, min, max bang CONG THUC Excel tham chieu Chi_tiet_5_lan, bam vao o la thay nguon."""
    detail, meta = R["detail"], R["meta"]
    wb = load_workbook(xlsx); ws, src = wb["Tong_hop"], wb["Chi_tiet_5_lan"]
    FONT, fill, thin = "Arial", PatternFill("solid", fgColor="DCE6F5"), Side(style="thin", color="BFBFBF")
    col_of = {src.cell(1, j).value: get_column_letter(j) for j in range(1, src.max_column + 1)}
    ws["A1"] = "TỔNG HỢP KẾT QUẢ 5 LẦN CHẠY"; ws["A1"].font = Font(name=FONT, size=13, bold=True)
    ws["A2"] = f"{len(meta)} ảnh, 5 lớp, {N_COMPONENTS} chiều; mỗi lần chia lại train/val/test và chọn lại siêu tham số trên val. Mọi số trong bảng là công thức tính từ sheet Chi_tiet_5_lan."
    ws["A2"].font = Font(name=FONT, size=10, italic=True, color="555555")
    METRICS = ["Acc train", "Acc val", "Acc test", "Precision macro test", "Recall macro test", "F1 macro test",
               "Acc test tự chụp", "Acc test RealWaste", "Loss train", "Loss val", "Thời gian tìm + huấn luyện (giây)"]
    STATS = [("Trung bình", "AVERAGE"), ("Độ lệch chuẩn", "STDEV"), ("Min", "MIN"), ("Max", "MAX")]
    r = 4
    for name in MODEL_NAMES:
        rows = detail.index[detail["Mô hình"] == name] + 2
        r0, r1 = rows.min(), rows.max()
        ws.cell(r, 1, name).font = Font(name=FONT, size=11, bold=True)
        ws.cell(r, 2, f"(dòng {r0} đến {r1} của Chi_tiet_5_lan)").font = Font(name=FONT, size=9, italic=True, color="777777"); r += 1
        ws.cell(r, 1, "Chỉ số")
        for j, (lbl, _) in enumerate(STATS): ws.cell(r, 2 + j, lbl)
        ws.cell(r, 2 + len(STATS), "Viết vào báo cáo (TB ± ĐLC)")
        for j in range(1, 3 + len(STATS)):
            c = ws.cell(r, j); c.font = Font(name=FONT, size=10, bold=True); c.fill = fill; c.border = Border(bottom=thin)
        r += 1
        for met in METRICS:
            rng_ = f"Chi_tiet_5_lan!{col_of[met]}{r0}:{col_of[met]}{r1}"
            ws.cell(r, 1, met).font = Font(name=FONT, size=10)
            for j, (_, fn) in enumerate(STATS):
                c = ws.cell(r, 2 + j, f"={fn}({rng_})"); c.font = Font(name=FONT, size=10); c.number_format = "0.0" if "Thời gian" in met else "0.0000"
            nd = 1 if "Thời gian" in met else 4
            ws.cell(r, 2 + len(STATS), f'=FIXED(AVERAGE({rng_}),{nd},TRUE)&" ± "&FIXED(STDEV({rng_}),{nd},TRUE)').font = Font(name=FONT, size=10)
            r += 1
        r += 1
    ws.cell(r, 1, "Accuracy test và siêu tham số được chọn ở từng lần chạy").font = Font(name=FONT, size=11, bold=True); r += 1
    ws.cell(r, 1, "Mô hình")
    for k in range(len(SEEDS)): ws.cell(r, 2 + k, f"Lần {k + 1} (seed {SEEDS[k]})")
    ws.cell(r, 2 + len(SEEDS), "Trung bình")
    for j in range(1, 3 + len(SEEDS)):
        c = ws.cell(r, j); c.font = Font(name=FONT, size=10, bold=True); c.fill = fill
    r += 1
    for name in MODEL_NAMES:
        rows = detail.index[detail["Mô hình"] == name] + 2
        ws.cell(r, 1, name).font = Font(name=FONT, size=10)
        for k, rr in enumerate(rows):
            c = ws.cell(r, 2 + k, f"=Chi_tiet_5_lan!{col_of['Acc test']}{rr}"); c.number_format = "0.0000"; c.font = Font(name=FONT, size=10)
            ws.cell(r + 1, 2 + k, f"=Chi_tiet_5_lan!{col_of['Siêu tham số chọn']}{rr}").font = Font(name=FONT, size=8, italic=True, color="555555")
        c = ws.cell(r, 2 + len(SEEDS), f"=AVERAGE({get_column_letter(2)}{r}:{get_column_letter(1 + len(SEEDS))}{r})")
        c.number_format = "0.0000"; c.font = Font(name=FONT, size=10, bold=True); r += 2
    ws.column_dimensions["A"].width = 26
    for j in range(2, 3 + max(len(STATS), len(SEEDS))): ws.column_dimensions[get_column_letter(j)].width = 22
    wb.save(xlsx)

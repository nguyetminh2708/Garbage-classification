# -*- coding: utf-8 -*-
"""
Phân loại rác: huấn luyện 3 mô hình ML cổ điển, khảo sát số chiều PCA và siêu tham số, chạy 5 lần, xuất Excel.

Chạy sau khi đã chạy tien_xu_ly_anh.py:
    python huan_luyen_mo_hinh.py

Kết quả (Excel, hình, mô hình) lưu vào thư mục ket_qua/.
"""


# ============================================================================
# Phân loại rác: 3 mô hình ML cổ điển, 2 khảo sát, 5 lần chạy
#
# Bài toán có 5 lớp của nhóm: giay, nhua, huu_co, kim_loai, khac (pin, thủy tinh, gốm, ...).
# Đầu vào là đặc trưng 512 chiều từ ResNet18 (features_512_raw.npz) và metadata.csv, do script tien_xu_ly_anh.py tạo ra.
#
# | # | Thuật toán | Họ | Biểu đồ loss |
# | --- | --- | --- | --- |
# | 1 | Softmax Regression (SGD) | Tuyến tính | Log loss train/val theo epoch |
# | 2 | SVM kernel RBF | Dựa trên lề (kernel) | Validation curve: hinge loss train/val theo C |
# | 3 | XGBoost (Gradient Boosting) | Ensemble cây (boosting) | Log loss train/val theo vòng boosting |
#
# Script chạy 3 phần, mỗi phần đều lặp 5 lần với 5 cách chia dữ liệu khác nhau (seed 0..4):
#
# Phần A. Khảo sát số chiều PCA: thử 32, 64, 128, 256, 512 (512 = không giảm chiều) với siêu tham số mặc định của từng mô hình,
#         để có cơ sở chọn số chiều dùng cho các phần sau.
# Phần B. Khảo sát siêu tham số: với số chiều đã chọn (N_COMPONENTS), thử MỌI tổ hợp trong lưới của từng mô hình.
#         Mỗi lần chạy chọn tổ hợp tốt nhất trên tập val; tổ hợp đó được đánh giá trên tập test của lần đó.
# Phần C. Tổng hợp 5 lần: trung bình ± độ lệch chuẩn, biểu đồ loss, ma trận nhầm lẫn, so sánh mô hình.
#
# Mọi con số (từng tổ hợp, từng lần, từng epoch) được ghi vào file Excel ket_qua_5_lan_chay.xlsx để làm dẫn chứng.
#
# Cách dùng: chạy python tien_xu_ly_anh.py trước, rồi chạy python huan_luyen_mo_hinh.py. Không cần GPU.
# Thời gian: khoảng 30 đến 60 phút trên CPU tùy máy (phần B thử 96 tổ hợp x 5 lần = 480 lượt huấn luyện). Có thể đặt RUN_PCA_SWEEP = False để bỏ phần A.
# ============================================================================

# ============================================================================
# 1. Cấu hình
# ============================================================================
import os as _os
try:
    BASE_DIR = _os.path.dirname(_os.path.abspath(__file__)).replace("\\", "/")   # thư mục chứa file script này
except NameError:
    BASE_DIR = _os.getcwd().replace("\\", "/")
FEAT_DIR = f"{BASE_DIR}/features"            # output của script tiền xử lý
OUT_DIR  = f"{BASE_DIR}/ket_qua"             # nơi lưu Excel, hình, mô hình

TOTAL_SAMPLES = 2250          # tổng số ảnh dùng (train + val + test); None = dùng hết
N_RUNS = 5                    # số lần chạy
SEEDS = [0, 1, 2, 3, 4]       # seed của từng lần chạy
SAMPLE_SEED = 42              # seed để chọn ảnh (cố định, 5 lần dùng chung một bộ ảnh)
VAL_SIZE, TEST_SIZE = 0.15, 0.15

# Phần A: khảo sát số chiều PCA
RUN_PCA_SWEEP = True
PCA_DIMS = [32, 64, 128, 256, 512]    # 512 = giữ nguyên, không PCA
N_COMPONENTS = 128                    # số chiều dùng cho phần B và C (chọn sau khi xem phần A)

# Phần B: lưới siêu tham số (thử MỌI tổ hợp, mỗi tổ hợp 5 lần)
MAX_EPOCHS = 200              # số epoch tối đa cho Softmax (SGD)
PATIENCE = 15                 # early stopping: dừng nếu loss val không giảm sau chừng này epoch
SOFTMAX_GRID = {"penalty": ["l2", "l1"], "alpha": [1e-5, 1e-4, 1e-3, 1e-2, 1e-1], "eta0": [1e-4, 1e-3, 1e-2]}   # 30 tổ hợp
SVM_GRID = {"kernel": ["linear", "rbf"], "C": [0.1, 1, 10, 100, 1000], "gamma": ["scale", 1e-4, 1e-3, 1e-2, 1e-1]}  # linear 5 + rbf 25 = 30 tổ hợp
XGB_GRID = {"max_depth": [3, 5, 7], "learning_rate": [0.05, 0.1, 0.3], "subsample": [0.7, 1.0], "colsample_bytree": [0.5, 1.0]}  # 36 tổ hợp
XGB_MAX_ROUNDS = 2000
XGB_EARLY_STOP = 50

# Siêu tham số mặc định dùng cho phần A (khảo sát PCA)
DEFAULT_HP = {"Softmax Regression": {"penalty": "l2", "alpha": 1e-3, "eta0": 1e-3},
              "SVM RBF": {"kernel": "rbf", "C": 10, "gamma": "scale"},
              "XGBoost": {"max_depth": 5, "learning_rate": 0.05, "subsample": 0.7, "colsample_bytree": 0.5}}

assert len(SEEDS) == N_RUNS


# ============================================================================
# 2. Đọc dữ liệu và chọn ảnh
# ============================================================================
import os, sys, copy, time, warnings, itertools

import matplotlib
matplotlib.use("Agg")   # lưu hình ra file, không mở cửa sổ
import numpy as np, pandas as pd, matplotlib.pyplot as plt
from sklearn.model_selection import train_test_split
from sklearn.preprocessing import StandardScaler
from sklearn.decomposition import PCA
from sklearn.linear_model import SGDClassifier
from sklearn.svm import SVC
from sklearn.metrics import (log_loss, hinge_loss, accuracy_score, f1_score, precision_score,
                             recall_score, precision_recall_fscore_support, confusion_matrix)
import xgboost as xgb
import joblib

os.makedirs(OUT_DIR, exist_ok=True)
raw = np.load(f"{FEAT_DIR}/features_512_raw.npz", allow_pickle=True)
meta_all = pd.read_csv(f"{FEAT_DIR}/metadata.csv")
F_all = raw["feats"]
assert len(meta_all) == len(F_all), "metadata.csv không khớp features_512_raw.npz, chạy lại script tiền xử lý"

CLASSES = ["giay", "nhua", "huu_co", "kim_loai", "khac"]
LABELS = np.arange(len(CLASSES))
MODEL_NAMES = ["Softmax Regression", "SVM RBF", "XGBoost"]
meta_all["y"] = meta_all.label.map({c: i for i, c in enumerate(CLASSES)})

def strat_split(index, test_size, strat, seed):
    """Chia phân tầng theo (lớp, nguồn); nhóm nào quá ít ảnh thì lùi về phân tầng theo lớp."""
    try:
        return train_test_split(index, test_size=test_size, stratify=strat, random_state=seed)
    except ValueError:
        warnings.warn("Có nhóm (lớp, nguồn) quá ít ảnh, chỉ phân tầng theo lớp.")
        return train_test_split(index, test_size=test_size, stratify=meta_all.loc[index, "label"], random_state=seed)

if TOTAL_SAMPLES is not None and len(meta_all) > TOTAL_SAMPLES:
    strat = meta_all.label + "|" + meta_all.source
    keep, _ = strat_split(meta_all.index, len(meta_all) - TOTAL_SAMPLES, strat, SAMPLE_SEED)
    meta = meta_all.loc[sorted(keep)].reset_index(drop=True)
    F = F_all[sorted(keep)]
else:
    if TOTAL_SAMPLES is not None and len(meta_all) < TOTAL_SAMPLES:
        warnings.warn(f"Chỉ có {len(meta_all)} ảnh, ít hơn {TOTAL_SAMPLES}. Dùng toàn bộ.")
    meta, F = meta_all.reset_index(drop=True), F_all
meta = meta.drop(columns=["split"], errors="ignore")
print(f"Dùng {len(meta)} ảnh, đặc trưng {F.shape[1]} chiều.")
tab = pd.crosstab(meta.label, meta.source, margins=True, margins_name="Tổng").reindex(CLASSES + ["Tổng"])
print(tab.to_string())


# ============================================================================
# 3. Hàm dùng chung
# ============================================================================
XGB_DEVICE = "cpu"
try:
    import torch
    XGB_DEVICE = "cuda" if torch.cuda.is_available() else "cpu"
except ImportError:
    pass

def style(ax):
    ax.spines[["top", "right"]].set_visible(False); ax.grid(alpha=0.25)

def hp_str(hp):
    return ", ".join(f"{k}={v}" for k, v in hp.items())

def prepare_run(seed, n_components):
    """Chia lại dữ liệu và fit lại scaler + PCA trên train của lần chạy này. n_components = 512 thì không PCA."""
    strat = meta.label + "|" + meta.source
    i_tr, i_tmp = strat_split(meta.index, VAL_SIZE + TEST_SIZE, strat, seed)
    i_va, i_te = strat_split(i_tmp, TEST_SIZE / (VAL_SIZE + TEST_SIZE), strat[i_tmp], seed)
    idx = {"train": np.sort(i_tr), "val": np.sort(i_va), "test": np.sort(i_te)}
    s1 = StandardScaler().fit(F[idx["train"]])
    if n_components >= F.shape[1]:
        pca, var_kept = None, 1.0
        Z = {s: s1.transform(F[idx[s]]) for s in idx}
    else:
        n_comp = min(n_components, len(idx["train"]))
        pca = PCA(n_components=n_comp, random_state=seed).fit(s1.transform(F[idx["train"]]))
        var_kept = pca.explained_variance_ratio_.sum()
        Z = {s: pca.transform(s1.transform(F[idx[s]])) for s in idx}
    s2 = StandardScaler().fit(Z["train"])
    X = {s: s2.transform(Z[s]) for s in idx}
    y = {s: meta.y.to_numpy()[idx[s]] for s in idx}
    return X, y, idx, {"scaler_512": s1, "pca": pca, "scaler_pca": s2}, var_kept

def fit_softmax(X, y, hp, seed):
    """Softmax Regression bằng SGD, theo dõi log loss từng epoch, early stopping theo loss val.
    Trả về (mô hình tốt nhất, lịch sử loss, epoch dừng)."""
    rng = np.random.default_rng(seed)
    clf = SGDClassifier(loss="log_loss", penalty=hp["penalty"], alpha=hp["alpha"], learning_rate="constant", eta0=hp["eta0"], random_state=seed)
    h = {"train_loss": [], "val_loss": [], "train_acc": [], "val_acc": []}
    best, best_loss, best_ep, wait = None, np.inf, 0, 0
    for ep in range(1, MAX_EPOCHS + 1):
        p = rng.permutation(len(y["train"]))
        clf.partial_fit(X["train"][p], y["train"][p], classes=LABELS)
        h["train_loss"].append(log_loss(y["train"], clf.predict_proba(X["train"]), labels=LABELS))
        h["val_loss"].append(log_loss(y["val"], clf.predict_proba(X["val"]), labels=LABELS))
        h["train_acc"].append(clf.score(X["train"], y["train"]))
        h["val_acc"].append(clf.score(X["val"], y["val"]))
        if h["val_loss"][-1] < best_loss - 1e-4:
            best, best_loss, best_ep, wait = copy.deepcopy(clf), h["val_loss"][-1], ep, 0
        else:
            wait += 1
            if wait >= PATIENCE:
                break
    return best, h, best_ep

def fit_svm(X, y, hp, seed):
    kw = {"gamma": hp["gamma"]} if hp["kernel"] == "rbf" else {}
    m = SVC(kernel=hp["kernel"], C=hp["C"], random_state=seed, **kw).fit(X["train"], y["train"])
    h = {"train_loss": [hinge_loss(y["train"], m.decision_function(X["train"]), labels=LABELS)],
         "val_loss": [hinge_loss(y["val"], m.decision_function(X["val"]), labels=LABELS)],
         "train_acc": [m.score(X["train"], y["train"])], "val_acc": [m.score(X["val"], y["val"])]}
    return m, h, None

def fit_xgb(X, y, hp, seed):
    m = xgb.XGBClassifier(objective="multi:softprob", n_estimators=XGB_MAX_ROUNDS, learning_rate=hp["learning_rate"],
                          max_depth=hp["max_depth"], subsample=hp["subsample"], colsample_bytree=hp["colsample_bytree"], min_child_weight=2,
                          reg_lambda=1.0, eval_metric="mlogloss", early_stopping_rounds=XGB_EARLY_STOP,
                          tree_method="hist", device=XGB_DEVICE, random_state=seed)
    m.fit(X["train"], y["train"], eval_set=[(X["train"], y["train"]), (X["val"], y["val"])], verbose=False)
    ev = m.evals_result(); bi = m.best_iteration
    h = {"train_loss": ev["validation_0"]["mlogloss"], "val_loss": ev["validation_1"]["mlogloss"],
         "train_acc": [m.score(X["train"], y["train"])], "val_acc": [m.score(X["val"], y["val"])]}
    return m, h, bi + 1

FIT = {"Softmax Regression": fit_softmax, "SVM RBF": fit_svm, "XGBoost": fit_xgb}
GRID = {"Softmax Regression": SOFTMAX_GRID, "SVM RBF": SVM_GRID, "XGBoost": XGB_GRID}
LOSS_NAME = {"Softmax Regression": "log loss", "SVM RBF": "hinge loss", "XGBoost": "log loss"}
# Tiêu chí chọn siêu tham số trên tập val: loss val thấp nhất (Softmax, XGBoost); accuracy val cao nhất, hòa thì loss val thấp hơn (SVM RBF)
def select_key(name, r):
    if name == "SVM RBF":
        return (-r["Acc val"], r["Loss val"])
    return (r["Loss val"], -r["Acc val"])

def fit_one(name, X, y, hp, seed):
    """Huấn luyện 1 mô hình với 1 bộ siêu tham số; trả về mô hình, lịch sử loss và dòng kết quả trên train/val."""
    t = time.time()
    m, h, be = FIT[name](X, y, hp, seed)
    k = (be - 1) if be else -1                      # vị trí điểm dừng trong lịch sử loss
    ka = k if len(h["train_acc"]) == len(h["train_loss"]) else -1   # XGBoost, SVM chỉ ghi accuracy của mô hình cuối
    r = {"Mô hình": name, "Siêu tham số": hp_str(hp), **hp, "Epoch/vòng dừng": be,
         "Loss train": h["train_loss"][k], "Loss val": h["val_loss"][k],
         "Acc train": h["train_acc"][ka], "Acc val": h["val_acc"][ka], "Thời gian (giây)": time.time() - t}
    return m, h, r

def grid_combos(grid):
    """Mọi tổ hợp của lưới. Riêng SVM kernel linear không dùng gamma nên gamma ghi "-" và bỏ các tổ hợp trùng."""
    keys = list(grid)
    out, seen = [], set()
    for vals in itertools.product(*[grid[k] for k in keys]):
        hp = dict(zip(keys, vals))
        if hp.get("kernel") == "linear":
            hp["gamma"] = "-"
        key = tuple(hp.items())
        if key not in seen:
            seen.add(key); out.append(hp)
    return out

def eval_test(m, X, y, src_test):
    p = m.predict(X["test"])
    out = {"Acc test": accuracy_score(y["test"], p),
           "Precision macro test": precision_score(y["test"], p, average="macro", zero_division=0),
           "Recall macro test": recall_score(y["test"], p, average="macro", zero_division=0),
           "F1 macro test": f1_score(y["test"], p, average="macro", zero_division=0),
           "Acc test tự chụp": accuracy_score(y["test"][src_test == "tu_chup"], p[src_test == "tu_chup"]) if (src_test == "tu_chup").any() else None,
           "Acc test RealWaste": accuracy_score(y["test"][src_test == "realwaste"], p[src_test == "realwaste"]) if (src_test == "realwaste").any() else None}
    return out, p


# ============================================================================
# 4. Phần A: khảo sát số chiều PCA
#
# Với mỗi số chiều trong PCA_DIMS, mỗi lần chạy chia lại dữ liệu, fit lại PCA và huấn luyện 3 mô hình
# với siêu tham số mặc định. Ghi accuracy val/test, phương sai giữ lại và thời gian.
# ============================================================================
pca_rows = []
if RUN_PCA_SWEEP:
    print("\n=== Phần A: khảo sát số chiều PCA ===")
    for dim in PCA_DIMS:
        t_dim = time.time()
        for run, seed in enumerate(SEEDS, start=1):
            X, y, idx, _, var_kept = prepare_run(seed, dim)
            src_test = meta.source.to_numpy()[idx["test"]]
            for name in MODEL_NAMES:
                m, h, r = fit_one(name, X, y, DEFAULT_HP[name], seed)
                te, _ = eval_test(m, X, y, src_test)
                pca_rows.append({"Số chiều": dim, "Mô hình": name, "Lần chạy": run, "Seed": seed,
                                 "PCA giữ phương sai": var_kept, "Siêu tham số": r["Siêu tham số"],
                                 "Epoch/vòng dừng": r["Epoch/vòng dừng"], "Loss train": r["Loss train"], "Loss val": r["Loss val"],
                                 "Acc train": r["Acc train"], "Acc val": r["Acc val"], "Acc test": te["Acc test"],
                                 "F1 macro test": te["F1 macro test"], "Thời gian (giây)": r["Thời gian (giây)"]})
        sub = pd.DataFrame(pca_rows); sub = sub[sub["Số chiều"] == dim]
        s = " | ".join(f"{n}: {sub[sub['Mô hình'] == n]['Acc val'].mean():.3f}" for n in MODEL_NAMES)
        print(f"  {dim:>3} chiều ({time.time() - t_dim:.0f}s)  acc val TB  {s}")
pca_df = pd.DataFrame(pca_rows)


# ============================================================================
# 5. Phần B: khảo sát siêu tham số và chọn mô hình từng lần chạy
#
# Với N_COMPONENTS chiều, mỗi lần chạy thử MỌI tổ hợp trong lưới của từng mô hình (ghi hết vào tuning_rows).
# Tổ hợp tốt nhất trên val của lần đó được giữ lại làm mô hình chính thức và đánh giá trên test (một lần duy nhất).
# ============================================================================
print(f"\n=== Phần B: khảo sát siêu tham số ({N_COMPONENTS} chiều) ===")
n_combo = {n: len(grid_combos(GRID[n])) for n in MODEL_NAMES}
print("Số tổ hợp thử mỗi lần chạy:", n_combo, "-> tổng", sum(n_combo.values()) * N_RUNS, "lượt huấn luyện")

detail_rows, class_rows, cm_rows, loss_rows, tuning_rows, data_rows = [], [], [], [], [], []
hists, saved_models = {}, {}

for run, seed in enumerate(SEEDS, start=1):
    t_run = time.time()
    X, y, idx, prep, var_kept = prepare_run(seed, N_COMPONENTS)
    src_test = meta.source.to_numpy()[idx["test"]]
    for s in idx:
        sub = meta.loc[idx[s]]
        for (lbl, src), n in sub.groupby(["label", "source"]).size().items():
            data_rows.append({"Lần chạy": run, "Seed": seed, "Tập": s, "Lớp": lbl, "Nguồn": src, "Số ảnh": n})

    models = {}
    for name in MODEL_NAMES:
        t_model = time.time()
        best = None
        for hp in grid_combos(GRID[name]):
            m, h, r = fit_one(name, X, y, hp, seed)
            tuning_rows.append({"Lần chạy": run, "Seed": seed, **r})
            if best is None or select_key(name, r) < select_key(name, best[2]):
                best = (m, h, r)
        m, h, r = best
        models[name] = m; hists[(run, name)] = (h, r["Epoch/vòng dừng"])
        te, p_test = eval_test(m, X, y, src_test)
        detail_rows.append({
            "Mô hình": name, "Lần chạy": run, "Seed": seed, "Siêu tham số chọn": r["Siêu tham số"],
            "Số ảnh train": len(y["train"]), "Số ảnh val": len(y["val"]), "Số ảnh test": len(y["test"]),
            "Số chiều": N_COMPONENTS, "PCA giữ phương sai": var_kept, "Epoch/vòng dừng": r["Epoch/vòng dừng"],
            "Loại loss": LOSS_NAME[name], "Loss train": r["Loss train"], "Loss val": r["Loss val"],
            "Acc train": r["Acc train"], "Acc val": r["Acc val"], **te,
            "Số tổ hợp đã thử": n_combo[name], "Thời gian tìm + huấn luyện (giây)": time.time() - t_model})
        pr, rc, f1, sp = precision_recall_fscore_support(y["test"], p_test, labels=LABELS, zero_division=0)
        for k, c in enumerate(CLASSES):
            class_rows.append({"Mô hình": name, "Lần chạy": run, "Lớp": c, "Precision": pr[k], "Recall": rc[k], "F1": f1[k], "Số ảnh test": sp[k]})
        cm = confusion_matrix(y["test"], p_test, labels=LABELS)
        for k, c in enumerate(CLASSES):
            cm_rows.append({"Mô hình": name, "Lần chạy": run, "Thực tế": c, **{f"Dự đoán {cc}": int(cm[k, j]) for j, cc in enumerate(CLASSES)}})
        if name != "SVM RBF":
            be = r["Epoch/vòng dừng"]
            for e in range(len(h["train_loss"])):
                loss_rows.append({"Mô hình": name, "Lần chạy": run, "Epoch/vòng": e + 1,
                                  "Loss train": h["train_loss"][e], "Loss val": h["val_loss"][e],
                                  "Được chọn (early stopping)": "x" if e + 1 == be else ""})
    saved_models[run] = {"prep": prep, "models": models}
    accs = " | ".join(f"{r['Mô hình']}: {r['Acc test']:.3f} ({r['Siêu tham số chọn']})" for r in detail_rows[-len(MODEL_NAMES):])
    print(f"Lần {run} (seed {seed}, {time.time() - t_run:.0f}s)  acc test  {accs}")

detail = pd.DataFrame(detail_rows)
detail["_o"] = detail["Mô hình"].map({n: i for i, n in enumerate(MODEL_NAMES)})
detail = detail.sort_values(["_o", "Lần chạy"]).drop(columns="_o").reset_index(drop=True)
tune = pd.DataFrame(tuning_rows)

# Tổng hợp từng tổ hợp qua 5 lần: trung bình ± độ lệch chuẩn trên val, và số lần tổ hợp đó được chọn
hp_summary = (tune.groupby(["Mô hình", "Siêu tham số"], sort=False)
              .agg(**{"Số lần": ("Acc val", "size"), "Acc train TB": ("Acc train", "mean"),
                      "Acc val TB": ("Acc val", "mean"), "Acc val ĐLC": ("Acc val", "std"),
                      "Loss val TB": ("Loss val", "mean"), "Loss val ĐLC": ("Loss val", "std"),
                      "Epoch/vòng dừng TB": ("Epoch/vòng dừng", "mean"), "Thời gian TB (giây)": ("Thời gian (giây)", "mean")})
              .reset_index())
chosen = detail.groupby(["Mô hình", "Siêu tham số chọn"]).size().rename("Số lần được chọn").reset_index().rename(columns={"Siêu tham số chọn": "Siêu tham số"})
hp_summary = hp_summary.merge(chosen, how="left", on=["Mô hình", "Siêu tham số"]).fillna({"Số lần được chọn": 0})
hp_summary["Số lần được chọn"] = hp_summary["Số lần được chọn"].astype(int)
hp_summary["_o"] = hp_summary["Mô hình"].map({n: i for i, n in enumerate(MODEL_NAMES)})
hp_summary = hp_summary.sort_values(["_o", "Acc val TB"], ascending=[True, False]).drop(columns="_o").reset_index(drop=True)


# ============================================================================
# 6. Phần C: kết quả trung bình 5 lần
# ============================================================================
METRICS = ["Acc train", "Acc val", "Acc test", "Precision macro test", "Recall macro test", "F1 macro test", "Acc test tự chụp", "Acc test RealWaste"]
g = detail.groupby("Mô hình", sort=False)
summary = pd.DataFrame({m: g[m].mean().map("{:.4f}".format) + " ± " + g[m].std().map("{:.4f}".format) for m in METRICS})
print("\n=== Kết quả trung bình 5 lần ===")
print(summary.to_string())


# ============================================================================
# 7. Biểu đồ (vẽ đơn giản bằng matplotlib mặc định: train màu xanh, val màu cam)
# ============================================================================
# 7a. Khảo sát PCA: accuracy val trung bình 5 lần theo số chiều, mỗi mô hình một đường
if len(pca_df):
    plt.figure(figsize=(7, 4))
    for name in MODEL_NAMES:
        mu = pca_df[pca_df["Mô hình"] == name].groupby("Số chiều")["Acc val"].mean()
        plt.plot(mu.index, mu.values, marker="o", label=name)
    plt.xscale("log", base=2); plt.xticks(PCA_DIMS, [str(d) for d in PCA_DIMS])
    plt.xlabel("Số chiều sau PCA (512 = không giảm chiều)"); plt.ylabel("Accuracy val (trung bình 5 lần)")
    plt.title("Khảo sát số chiều PCA"); plt.legend(); plt.grid(alpha=0.3)
    plt.tight_layout(); plt.savefig(f"{OUT_DIR}/khao_sat_pca.png", dpi=120); plt.close()

# 7b. Ảnh hưởng của từng siêu tham số (chỉ ghi ra Excel, không vẽ): với mỗi giá trị của một tham số,
#     lấy accuracy val TB tốt nhất khi các tham số còn lại tự do
for name in MODEL_NAMES:
    for k in GRID[name]:
        if k not in hp_summary: hp_summary[k] = None
    for i in hp_summary.index[hp_summary["Mô hình"] == name]:
        for part in hp_summary.at[i, "Siêu tham số"].split(", "):
            k, v = part.split("=")
            hp_summary.at[i, k] = v
tune["gamma"] = tune["gamma"].astype(str)
marg_rows = []
for name in MODEL_NAMES:
    sub = hp_summary[hp_summary["Mô hình"] == name]
    for k in GRID[name]:
        vals = [str(v) for v in GRID[name][k]] + (["-"] if k == "gamma" else [])
        for v in vals:
            ss = sub[sub[k].astype(str) == v]
            if not len(ss): continue
            marg_rows.append({"Mô hình": name, "Tham số": k, "Giá trị": v, "Số tổ hợp": len(ss),
                              "Acc val TB tốt nhất": ss["Acc val TB"].max(), "Acc val TB trung bình": ss["Acc val TB"].mean(),
                              "Acc train TB (tại tổ hợp tốt nhất)": ss.loc[ss["Acc val TB"].idxmax(), "Acc train TB"]})
marg = pd.DataFrame(marg_rows)

# 7c. Biểu đồ loss theo epoch / vòng của mô hình được chọn ở 5 lần chạy (mỗi lần một đường; chấm = điểm early stopping)
def plot_runs(name, xlabel, ylabel, fname):
    plt.figure(figsize=(7, 4))
    for r in range(1, N_RUNS + 1):
        h, be = hists[(r, name)]
        e = np.arange(1, len(h["train_loss"]) + 1)
        plt.plot(e, h["train_loss"], color="C0", label="train" if r == 1 else None)
        plt.plot(e, h["val_loss"], color="C1", label="val" if r == 1 else None)
        plt.plot(be, h["val_loss"][be - 1], "ko", ms=5, label="điểm dừng (early stopping)" if r == 1 else None)
    plt.xlabel(xlabel); plt.ylabel(ylabel); plt.title(f"{name}: loss train và val qua 5 lần chạy")
    plt.legend(); plt.grid(alpha=0.3)
    plt.tight_layout(); plt.savefig(f"{OUT_DIR}/{fname}", dpi=120); plt.close()

plot_runs("Softmax Regression", "Epoch", "Log loss", "loss_softmax_5_lan.png")
plot_runs("XGBoost", "Vòng boosting (số cây)", "Log loss", "loss_xgboost_5_lan.png")

# SVM RBF không lặp theo epoch: validation curve hinge loss theo C, tại gamma được chọn của từng lần
rbf = tune[tune["Mô hình"] == "SVM RBF"]
plt.figure(figsize=(7, 4))
for run in range(1, N_RUNS + 1):
    chosen_hp = detail[(detail["Mô hình"] == "SVM RBF") & (detail["Lần chạy"] == run)]["Siêu tham số chọn"].iloc[0]
    g_best = chosen_hp.split("gamma=")[1]
    sub = rbf[(rbf["Lần chạy"] == run) & (rbf.gamma == g_best)].sort_values("C")
    plt.plot(sub.C, sub["Loss train"], "o-", color="C0", label="train" if run == 1 else None)
    plt.plot(sub.C, sub["Loss val"], "o-", color="C1", label="val" if run == 1 else None)
plt.xscale("log"); plt.xlabel("C"); plt.ylabel("Hinge loss")
plt.title("SVM RBF: hinge loss theo C qua 5 lần chạy"); plt.legend(); plt.grid(alpha=0.3)
plt.tight_layout(); plt.savefig(f"{OUT_DIR}/loss_svm_rbf_5_lan.png", dpi=120); plt.close()

# 7d. Ma trận nhầm lẫn (cộng dồn 5 lần) và so sánh mô hình
from sklearn.metrics import ConfusionMatrixDisplay
cmdf = pd.DataFrame(cm_rows)
fig, axes = plt.subplots(1, 3, figsize=(15, 4.6))
for ax, name in zip(axes, MODEL_NAMES):
    cm = cmdf[cmdf["Mô hình"] == name].groupby("Thực tế", sort=False)[[f"Dự đoán {c}" for c in CLASSES]].sum().reindex(CLASSES).to_numpy()
    ConfusionMatrixDisplay(cm, display_labels=CLASSES).plot(ax=ax, colorbar=False, xticks_rotation=30)
    ax.set_title(name); ax.set_xlabel("Dự đoán"); ax.set_ylabel("Thực tế")
plt.tight_layout(); plt.savefig(f"{OUT_DIR}/confusion_matrix_5_lan.png", dpi=120); plt.close()

means = g[["Acc test", "F1 macro test"]].mean().reindex(MODEL_NAMES)
ax = means.plot.bar(figsize=(7, 4), rot=0)
for cont in ax.containers:
    ax.bar_label(cont, fmt="%.3f", fontsize=9)
ax.set_ylim(0, 1); ax.set_xlabel(""); ax.set_title("Accuracy và Macro F1 trên tập test (trung bình 5 lần)")
ax.legend(["Accuracy test", "Macro F1 test"]); ax.grid(axis="y", alpha=0.3)
plt.tight_layout(); plt.savefig(f"{OUT_DIR}/so_sanh_mo_hinh_5_lan.png", dpi=120); plt.close()


# ============================================================================
# 8. Xuất file Excel
#
# File ket_qua_5_lan_chay.xlsx gồm các sheet:
#
# | Sheet | Nội dung |
# | --- | --- |
# | Tong_hop | Trung bình, độ lệch chuẩn, min, max của từng mô hình qua 5 lần. Tính bằng công thức Excel từ sheet Chi_tiet_5_lan |
# | Chi_tiet_5_lan | Mỗi dòng là 1 mô hình trong 1 lần chạy: seed, siêu tham số chọn, epoch dừng, loss, accuracy, precision, recall, F1, thời gian |
# | Khao_sat_PCA | Phần A: từng số chiều, từng mô hình, từng lần chạy |
# | Khao_sat_PCA_TB | Phần A tổng hợp: trung bình ± ĐLC theo số chiều và mô hình |
# | Sieu_tham_so_TB | Phần B tổng hợp: mỗi tổ hợp siêu tham số, trung bình ± ĐLC trên val qua 5 lần, số lần được chọn |
# | Anh_huong_tung_tham_so | Phần B: với từng giá trị của từng siêu tham số, accuracy val TB tốt nhất và trung bình khi các tham số còn lại tự do |
# | Sieu_tham_so_chi_tiet | Phần B: mọi tổ hợp đã thử trong từng lần chạy |
# | Theo_lop | Precision, recall, F1 của từng lớp, từng lần chạy |
# | Ma_tran_nham_lan | Ma trận nhầm lẫn trên test của từng lần chạy |
# | Loss_theo_epoch | Loss train/val ở từng epoch / vòng của mô hình được chọn, đánh dấu điểm early stopping |
# | Du_lieu_tung_lan | Số ảnh theo tập, lớp, nguồn của từng lần chia |
# | Cau_hinh | Các tham số cấu hình đã dùng, phiên bản thư viện, máy |
# ============================================================================
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

XLSX = f"{OUT_DIR}/ket_qua_5_lan_chay.xlsx"
config = pd.DataFrame([
    ("Tổng số ảnh dùng", len(meta)), ("Số lần chạy", N_RUNS), ("Seed các lần", ", ".join(map(str, SEEDS))),
    ("Seed chọn ảnh", SAMPLE_SEED), ("Tỉ lệ train / val / test", f"{1-VAL_SIZE-TEST_SIZE:.0%} / {VAL_SIZE:.0%} / {TEST_SIZE:.0%}"),
    ("Trích đặc trưng", "ResNet18 (ImageNet), 512 chiều"),
    ("Phần A: số chiều PCA thử", str(PCA_DIMS) if RUN_PCA_SWEEP else "bỏ qua"),
    ("Phần A: siêu tham số mặc định", "; ".join(f"{n}: {hp_str(h)}" for n, h in DEFAULT_HP.items())),
    ("Số chiều dùng cho phần B, C", f"StandardScaler + PCA {N_COMPONENTS} chiều, fit trên train từng lần"),
    ("Softmax: lưới", str(SOFTMAX_GRID)), ("Softmax: epoch tối đa / patience", f"{MAX_EPOCHS} / {PATIENCE}"),
    ("SVM RBF: lưới", str(SVM_GRID)),
    ("XGBoost: lưới", str(XGB_GRID)), ("XGBoost: vòng tối đa / early stopping", f"{XGB_MAX_ROUNDS} / {XGB_EARLY_STOP}"),
    ("XGBoost: cố định", "min_child_weight 2, reg_lambda 1"),
    ("Số tổ hợp thử mỗi lần chạy", str(n_combo)),
    ("Tiêu chí chọn siêu tham số", "Loss val thấp nhất (Softmax, XGBoost); accuracy val cao nhất, hòa thì loss val thấp hơn (SVM RBF)"),
    ("Ngày chạy", time.strftime("%Y-%m-%d %H:%M")),
    ("Python", sys.version.split()[0]),
    ("scikit-learn / XGBoost", f"{__import__('sklearn').__version__} / {xgb.__version__}"),
    ("Máy", f"{__import__('platform').platform()}; CPU {__import__('platform').processor() or 'không rõ'}; GPU {'có' if XGB_DEVICE == 'cuda' else 'không'}"),
], columns=["Tham số", "Giá trị"])

pca_tb = pd.DataFrame()
if len(pca_df):
    gp = pca_df.groupby(["Mô hình", "Số chiều"], sort=False)
    pca_tb = gp.agg(**{"PCA giữ phương sai": ("PCA giữ phương sai", "mean"), "Acc train TB": ("Acc train", "mean"),
                       "Acc val TB": ("Acc val", "mean"), "Acc val ĐLC": ("Acc val", "std"),
                       "Acc test TB": ("Acc test", "mean"), "Acc test ĐLC": ("Acc test", "std"),
                       "F1 test TB": ("F1 macro test", "mean"), "Thời gian TB (giây)": ("Thời gian (giây)", "mean")}).reset_index()

hp_cols = ["Mô hình", "Siêu tham số", "Số lần", "Acc train TB", "Acc val TB", "Acc val ĐLC", "Loss val TB", "Loss val ĐLC", "Epoch/vòng dừng TB", "Thời gian TB (giây)", "Số lần được chọn"]
tune_cols = ["Lần chạy", "Seed", "Mô hình", "Siêu tham số", "Epoch/vòng dừng", "Loss train", "Loss val", "Acc train", "Acc val", "Thời gian (giây)"]

with pd.ExcelWriter(XLSX, engine="openpyxl") as w:
    pd.DataFrame().to_excel(w, sheet_name="Tong_hop")
    detail.to_excel(w, sheet_name="Chi_tiet_5_lan", index=False)
    pca_df.to_excel(w, sheet_name="Khao_sat_PCA", index=False)
    pca_tb.to_excel(w, sheet_name="Khao_sat_PCA_TB", index=False)
    hp_summary[hp_cols].to_excel(w, sheet_name="Sieu_tham_so_TB", index=False)
    marg.to_excel(w, sheet_name="Anh_huong_tung_tham_so", index=False)
    tune[tune_cols].to_excel(w, sheet_name="Sieu_tham_so_chi_tiet", index=False)
    pd.DataFrame(class_rows).to_excel(w, sheet_name="Theo_lop", index=False)
    cmdf.to_excel(w, sheet_name="Ma_tran_nham_lan", index=False)
    pd.DataFrame(loss_rows).to_excel(w, sheet_name="Loss_theo_epoch", index=False)
    pd.DataFrame(data_rows).to_excel(w, sheet_name="Du_lieu_tung_lan", index=False)
    config.to_excel(w, sheet_name="Cau_hinh", index=False)

wb = load_workbook(XLSX)
FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="DCE6F5")
thin = Side(style="thin", color="BFBFBF")

# Định dạng chung cho các sheet dữ liệu
for ws in wb.worksheets[1:]:
    for row in ws.iter_rows():
        for c in row:
            c.font = Font(name=FONT, size=10, bold=(c.row == 1))
            if c.row == 1:
                c.fill = HEAD_FILL; c.alignment = Alignment(wrap_text=True, vertical="center")
    for col in ws.columns:
        head = str(col[0].value or "")
        width = max(len(head), *(len(str(c.value)) if c.value is not None else 0 for c in col[1:200]), 6)
        ws.column_dimensions[col[0].column_letter].width = min(width + 2, 45)
        if any(k in head for k in ["Acc", "Precision", "Recall", "F1", "PCA giữ", "Loss"]):
            for c in col[1:]: c.number_format = "0.0000"
        if "Thời gian" in head or "dừng TB" in head:
            for c in col[1:]: c.number_format = "0.0"
    ws.freeze_panes = "A2"

# Sheet Tong_hop: công thức tham chiếu Chi_tiet_5_lan
ws = wb["Tong_hop"]
src = wb["Chi_tiet_5_lan"]
col_of = {src.cell(1, j).value: get_column_letter(j) for j in range(1, src.max_column + 1)}
ws["A1"] = "TỔNG HỢP KẾT QUẢ 5 LẦN CHẠY"
ws["A2"] = f"{len(meta)} ảnh, 5 lớp, {N_COMPONENTS} chiều; mỗi lần chia lại train/val/test với seed khác và chọn lại siêu tham số trên val. Mọi số trong bảng là công thức tính từ sheet Chi_tiet_5_lan."
ws["A1"].font = Font(name=FONT, size=13, bold=True); ws["A2"].font = Font(name=FONT, size=10, italic=True, color="555555")

SUM_METRICS = ["Acc train", "Acc val", "Acc test", "Precision macro test", "Recall macro test", "F1 macro test", "Acc test tự chụp", "Acc test RealWaste", "Loss train", "Loss val", "Thời gian tìm + huấn luyện (giây)"]
STATS = [("Trung bình", "AVERAGE"), ("Độ lệch chuẩn", "STDEV"), ("Min", "MIN"), ("Max", "MAX")]
r = 4
for name in MODEL_NAMES:
    rows = detail.index[detail["Mô hình"] == name] + 2     # dòng trong Excel (có header)
    r0, r1 = rows.min(), rows.max()
    ws.cell(r, 1, name).font = Font(name=FONT, size=11, bold=True)
    ws.cell(r, 2, f"(dòng {r0} đến {r1} của Chi_tiet_5_lan)").font = Font(name=FONT, size=9, italic=True, color="777777")
    r += 1
    ws.cell(r, 1, "Chỉ số")
    for j, (lbl, _) in enumerate(STATS): ws.cell(r, 2 + j, lbl)
    ws.cell(r, 2 + len(STATS), "Viết vào báo cáo (TB ± ĐLC)")
    for j in range(1, 3 + len(STATS)):
        c = ws.cell(r, j); c.font = Font(name=FONT, size=10, bold=True); c.fill = HEAD_FILL; c.border = Border(bottom=thin)
    r += 1
    for met in SUM_METRICS:
        rng_ = f"Chi_tiet_5_lan!{col_of[met]}{r0}:{col_of[met]}{r1}"
        ws.cell(r, 1, met).font = Font(name=FONT, size=10)
        for j, (_, fn) in enumerate(STATS):
            c = ws.cell(r, 2 + j, f"={fn}({rng_})")
            c.font = Font(name=FONT, size=10); c.number_format = "0.0" if "Thời gian" in met else "0.0000"
        nd = 1 if "Thời gian" in met else 4   # FIXED tự dùng dấu thập phân theo ngôn ngữ máy
        c = ws.cell(r, 2 + len(STATS), f'=FIXED(AVERAGE({rng_}),{nd},TRUE)&" ± "&FIXED(STDEV({rng_}),{nd},TRUE)')
        c.font = Font(name=FONT, size=10)
        r += 1
    r += 1

# Bảng so sánh nhanh accuracy test từng lần và siêu tham số được chọn
ws.cell(r, 1, "Accuracy test và siêu tham số được chọn ở từng lần chạy").font = Font(name=FONT, size=11, bold=True); r += 1
ws.cell(r, 1, "Mô hình")
for k in range(N_RUNS): ws.cell(r, 2 + k, f"Lần {k + 1} (seed {SEEDS[k]})")
ws.cell(r, 2 + N_RUNS, "Trung bình")
for j in range(1, 3 + N_RUNS):
    c = ws.cell(r, j); c.font = Font(name=FONT, size=10, bold=True); c.fill = HEAD_FILL
r += 1
for name in MODEL_NAMES:
    rows = detail.index[detail["Mô hình"] == name] + 2
    ws.cell(r, 1, name).font = Font(name=FONT, size=10)
    for k, rr in enumerate(rows):
        c = ws.cell(r, 2 + k, f"=Chi_tiet_5_lan!{col_of['Acc test']}{rr}"); c.number_format = "0.0000"; c.font = Font(name=FONT, size=10)
        c = ws.cell(r + 1, 2 + k, f"=Chi_tiet_5_lan!{col_of['Siêu tham số chọn']}{rr}"); c.font = Font(name=FONT, size=8, italic=True, color="555555")
    c = ws.cell(r, 2 + N_RUNS, f"=AVERAGE({get_column_letter(2)}{r}:{get_column_letter(1 + N_RUNS)}{r})")
    c.number_format = "0.0000"; c.font = Font(name=FONT, size=10, bold=True)
    r += 2

ws.column_dimensions["A"].width = 26
for j in range(2, 3 + max(len(STATS), N_RUNS)):
    ws.column_dimensions[get_column_letter(j)].width = 22
ws.column_dimensions[get_column_letter(2 + len(STATS))].width = 26
wb.save(XLSX)
print("Đã ghi", XLSX)


# ============================================================================
# 9. Lưu mô hình
# ============================================================================
joblib.dump({"runs": saved_models, "classes": CLASSES, "seeds": SEEDS, "n_components": N_COMPONENTS}, f"{OUT_DIR}/mo_hinh_5_lan.joblib")
print("Đã lưu vào", OUT_DIR)
for f in sorted(os.listdir(OUT_DIR)): print("  ", f)


# ============================================================================
# 10. Gợi ý viết báo cáo
#
# Kết quả chính: cột "Viết vào báo cáo (TB ± ĐLC)" ở sheet Tong_hop.
# Vì sao chọn 128 chiều: sheet Khao_sat_PCA_TB và hình khao_sat_pca.png.
# Vì sao chọn siêu tham số đó: sheet Sieu_tham_so_TB (tổ hợp nào tốt nhất trên val, được chọn mấy lần trong 5 lần) và sheet Anh_huong_tung_tham_so.
# Hội tụ và không overfit: hình loss_*_5_lan.png; chấm tròn là điểm early stopping, sau đó loss val đi ngang hoặc tăng.
# Khi được hỏi về một lần chạy cụ thể: Chi_tiet_5_lan, Loss_theo_epoch, Sieu_tham_so_chi_tiet.
# ============================================================================
